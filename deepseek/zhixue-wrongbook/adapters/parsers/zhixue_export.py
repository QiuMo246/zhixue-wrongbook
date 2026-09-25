"""通道 A：官方导出文件解析（一等公民，合规降级通道）。

为什么这条通道重要（架构文档第 3 节）：
  1. 完全合规——用户自己从 APP 导出，工具只解析文件，不碰账号
  2. 接口全挂时的 fallback，链路照跑
  3. 官方数据可用来交叉校验 API 通道的结果

支持的输入：.pdf / .docx / .txt / .md
不支持的：扫描版 PDF（纯图片、无文本层）—— 这种情况会明确报
「未提取到文本」，并建议改用 OCR，而不是返回一堆空题。

诚实声明：PDF 的排版千变万化，本解析器基于「编号 + 标签」的通用启发式，
**不保证对所有学校的导出格式都准确**。所以每题都带 parse_confidence，
低置信度的题会被统计自动排除，并出现在复核队列里。
"""

from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path

from adapters.parsers.normalize import (estimate_parse_confidence, make_id,
                                        normalize_question)
from core.models import WrongQuestion

EXPORT_FORMAT_VERSION = "zhixue-export/heuristic-1"

# HTML → 带换行的纯文本（供题型判定用）。
# 块级标签换行、行内标签删掉，避免把「下列<span>的</span>是」拆断。
_BLOCK_END_RE = re.compile(r"<\s*/?\s*(?:p|div|li|tr|br|h[1-6])\b[^>]*>", re.I)
_ANY_TAG_RE = re.compile(r"<[^>]+>")


def _html_to_lines(html: str | None) -> str:
    s = _BLOCK_END_RE.sub("\n", html or "")
    return _ANY_TAG_RE.sub("", s)

# 平台的填空下划线：`<u class="blank blank1 _blank_1">`。
# 剥标签后下划线就没了，所以判填空题必须回头看原始 HTML。
_BLANK_MARK_RE = re.compile(r'class\s*=\s*"[^"]*\bblank\b', re.I)

# 题目起始：行首的 "1." "1、" "1．" "1)" 等
Q_START_RE = re.compile(r"^\s*(\d{1,3})\s*[.．、)）]\s*(?=\S)")
# 标签行
LABEL_RE = re.compile(
    r"^\s*(?:【?\s*(?P<tag>答案|参考答案|标准答案|解析|详解|解答|学生答案|我的答案|"
    r"作答|订正|错因|知识点|考点|难度|得分|满分)\s*】?)\s*[:：]?\s*(?P<rest>.*)$")
SCORE_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(?:/|分\s*/\s*|／)\s*(\d+(?:\.\d+)?)")
DIFF_RE = re.compile(r"(?:难度)\s*[:：]?\s*(0?\.\d+|1(?:\.0+)?)")
RATE_RE = re.compile(r"(?:班级得分率|得分率)\s*[:：]?\s*(0?\.\d+|1(?:\.0+)?|100(?:\.0+)?\s*%)")

TAG_ALIAS = {
    "答案": "standard", "参考答案": "standard", "标准答案": "standard",
    "解析": "analysis", "详解": "analysis", "解答": "analysis",
    "学生答案": "student", "我的答案": "student", "作答": "student",
    "订正": "correction", "错因": "error_hint", "知识点": "kp_hint",
    "考点": "kp_hint", "难度": "difficulty", "得分": "score", "满分": "score",
}


# ---------------------------------------------------------------------------
# 文本抽取
# ---------------------------------------------------------------------------
def extract_text(path: Path) -> tuple[str, str]:
    """返回 (文本, 提取方式说明)。"""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("未安装 pypdf，无法解析 PDF。"
                               "执行：.venv/Scripts/python -m pip install pypdf") from exc
        reader = PdfReader(str(path))
        pages = [(p.extract_text() or "") for p in reader.pages]
        text = "\n".join(pages)
        if not text.strip():
            raise RuntimeError(
                f"PDF 未提取到任何文本（共 {len(pages)} 页）。"
                "这通常是扫描版/纯图片 PDF。请改用带文本层的导出文件，"
                "或先做 OCR 再导入——本工具不做 OCR，也不会伪造结果。")
        return text, f"pypdf,{len(pages)}页"

    if suffix == ".docx":
        try:
            import docx
        except ImportError as exc:
            raise RuntimeError("未安装 python-docx。执行："
                               ".venv/Scripts/python -m pip install python-docx") from exc
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for tbl in d.tables:
            for row in tbl.rows:
                parts.append("\t".join(c.text for c in row.cells))
        return "\n".join(parts), "python-docx"

    if suffix in (".txt", ".md", ".markdown"):
        return path.read_text(encoding="utf-8", errors="replace"), "plain-text"

    raise RuntimeError(f"不支持的文件类型 {suffix}。支持：pdf / docx / txt / md")


# ---------------------------------------------------------------------------
# 结构化
# ---------------------------------------------------------------------------
def split_questions(text: str) -> list[dict]:
    """按题号切块，再按标签把块内的行归类。"""
    lines = text.splitlines()
    blocks: list[tuple[int, list[str]]] = []
    current: list[str] | None = None
    current_no = 0
    for ln in lines:
        m = Q_START_RE.match(ln)
        # 只有「看起来像题干开始」才切：题号后的内容不能太短，
        # 否则像 "1. A" 这种选项行会被误当成新题。
        if m and len(ln.strip()) > len(m.group(0)) + 2:
            if current is not None:
                blocks.append((current_no, current))
            current_no = int(m.group(1))
            current = [ln[m.end():].strip()]
        elif current is not None:
            current.append(ln)
    if current is not None:
        blocks.append((current_no, current))

    out = []
    for no, body in blocks:
        out.append(_parse_block(no, body))
    return out


def _parse_block(no: int, body: list[str]) -> dict:
    fields: dict[str, list[str]] = {"stem": []}
    cur = "stem"
    for ln in body:
        m = LABEL_RE.match(ln)
        if m:
            tag = TAG_ALIAS.get(m.group("tag"))
            if tag:
                if tag in ("difficulty", "score", "kp_hint", "error_hint"):
                    fields.setdefault(tag, []).append(ln.strip())
                    cur = "stem"  # 这类单行标签之后回到题干/答案流
                    continue
                cur = tag
                fields.setdefault(cur, [])
                rest = (m.group("rest") or "").strip()
                if rest:
                    fields[cur].append(rest)
                continue
        fields.setdefault(cur, []).append(ln)

    def join(key: str) -> str:
        return "\n".join(x for x in fields.get(key, []) if x.strip()).strip()

    stem = join("stem")
    standard = join("standard")
    analysis = join("analysis")
    student = join("student")
    correction = join("correction")
    difficulty = None
    rate = None
    got = full = None
    if fields.get("difficulty"):
        m = DIFF_RE.search(fields["difficulty"][0])
        if m:
            difficulty = float(m.group(1))
    for line in fields.get("score", []) + body:
        m = SCORE_RE.search(line)
        if m:
            got, full = float(m.group(1)), float(m.group(2))
            break
    for line in body:
        m = RATE_RE.search(line)
        if m:
            raw = m.group(1)
            rate = float(raw.rstrip("%")) / 100 if raw.endswith("%") else float(raw)
            break

    return {
        "no": no,
        "stem": stem,
        "standard": standard,
        "analysis": analysis,
        "student": student,
        "correction": correction,
        "difficulty": difficulty,
        "class_score_rate": rate,
        "score_got": got,
        "score_full": full,
    }


def _to_html(text: str) -> str:
    """纯文本 → 极简 HTML（保留换行）。图片在文本导出里不存在。"""
    if not text:
        return ""
    escaped = (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
    return "<p>" + escaped.replace("\n", "<br>") + "</p>"


def parse_file(path: str | Path, subject: str, exam_name: str,
               exam_date: str | None = None, grade: str | None = None,
               include_correction_as_student: bool = True) -> tuple[list[WrongQuestion], dict]:
    """解析导出文件 → WrongQuestion 列表 + 解析报告。"""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在：{path}")
    text, how = extract_text(path)
    blocks = split_questions(text)
    if not blocks:
        raise RuntimeError(
            "未识别出任何题目。可能原因：导出文件排版与通用启发式不匹配，"
            "或内容里没有「1. / 2.」这样的题号。请把文件发我，我按实际排版加规则——"
            "不会假装解析成功。")

    questions: list[WrongQuestion] = []
    stats = {"total_blocks": len(blocks), "with_standard": 0, "with_analysis": 0,
             "with_student": 0, "with_score": 0, "with_difficulty": 0,
             "low_confidence": 0}
    for b in blocks:
        student_text = b["student"] or (b["correction"] if include_correction_as_student else "")
        conf = estimate_parse_confidence(
            has_stem=bool(b["stem"]),
            has_standard=bool(b["standard"]),
            has_student=bool(student_text),
            has_analysis=bool(b["analysis"]),
            has_score=b["score_got"] is not None,
            has_difficulty=b["difficulty"] is not None,
        )
        raw = {
            "subject": subject,
            "exam_name": exam_name,
            "exam_date": exam_date,
            "grade": grade,
            "stem_html": _to_html(b["stem"]),
            "qtype": _guess_qtype(b["stem"], b["standard"]),
            "standard": b["standard"] or None,
            "analysis_html": _to_html(b["analysis"]) or None,
            "student_text": student_text or None,
            "difficulty": b["difficulty"],
            # 导出文件里的「难度」是文本解析出来的，刻度同样未经核实
            "difficulty_scale": ("0-1" if (b["difficulty"] is not None
                                           and 0.0 <= b["difficulty"] <= 1.0)
                                 else "unknown"),
            "class_score_rate": b["class_score_rate"],
            "score_got": b["score_got"],
            "score_full": b["score_full"],
            "parse_confidence": conf,
        }
        q = normalize_question(raw, source="export",
                               source_version=EXPORT_FORMAT_VERSION,
                               seq=b["no"],
                               fetched_at=datetime.now(timezone.utc))
        questions.append(q)
        stats["with_standard"] += bool(b["standard"])
        stats["with_analysis"] += bool(b["analysis"])
        stats["with_student"] += bool(student_text)
        stats["with_score"] += b["score_got"] is not None
        stats["with_difficulty"] += b["difficulty"] is not None
        stats["low_confidence"] += conf < 0.5

    stats["extract_method"] = how
    stats["file"] = str(path)
    stats["note"] = ("解析基于通用启发式，不保证与所有学校导出的排版匹配；"
                     "低置信度题目会被统计自动排除并进入复核队列。")
    return questions, stats


def _guess_qtype(stem: str, standard: str) -> str:
    """题型粗判。判不出来就给「解答题」——它是默认兜底，不是猜测的结论。

    匹配前先把 HTML 还原成带换行的纯文本：API 通道给的是 HTML，
    选项常写成 `<p>A. 甲</p><p>B. 乙</p>`，不剥标签的话
    `^\\s*[A-D][.．、]` 一条都匹配不上（行首是 `<p>` 而不是 `A`）。
    块级标签换成换行、行内标签直接删掉，这样「下列<span>的</span>是」
    也不会被拆断。

    **2026-09-24 重排判定顺序**（真实数据 27 道题暴露的两类系统性误判）：

    缺陷 1：探究题里的小问自带 A/B/C 选项，原来的「先看选项」把整题判成
      选择题。实测 4 道中招——
        · 物理「利用如图甲所示的实验装置观察水的沸腾：(1)…(5)」→ 单选题
        · 物理「请按要求回答：(1)…(5)」（声音探究）→ 多选题
        · 物理「探究平面镜成像特点」→ 多选题
        · 物理「探究凸透镜成像规律」→ 多选题
      后果不只是标签难看：`core/validate.py` 的练习题闸门要求
      「生成题题型 == 原题题型」，原题标错，生成的同类题就跟着错。
      修法：**「实验/探究 + ≥2 个小问」优先于选项判定**。

    缺陷 2：平台的填空下划线写成 `<u class="blank …">`，`html_to_text`
      把标签剥掉后下划线消失，`填空|____` 一条都匹配不上，填空题
      全被兜底成解答题。实测 5 道中招（数学 2 道、物理 3 道）。
      修法：**同时检查原始 HTML 里的 blank 类名**。

    注意「有选项就不算填空」这条护栏：英语的「完形填空」题干里带
    「填空」二字、又带选项，不加护栏会被判成填空题。
    """
    raw = stem or ""
    s = _html_to_lines(raw)
    has_options = bool(re.search(r"^\s*[A-D][.．、]", s, re.M))
    n_sub = len(set(re.findall(r"[（(]\s*([1-9])\s*[）)]", s)))

    # 1) 实验探究题：带「实验/探究」且有小问。必须排在选项判定之前。
    if n_sub >= 2 and re.search(r"实验|探究", s):
        return "实验探究题"
    # 2) 作图题：必须排在填空题之前——作图题的小问里常有填空横线。
    if re.search(r"作图|画出", s):
        return "作图题"
    # 3) 证明题
    if re.search(r"证明|求证", s):
        return "证明题"
    # 4) 填空题（有选项就不算，避免把完形填空/选词填空吃进来）
    if not has_options and (_BLANK_MARK_RE.search(raw) or re.search(r"填空|_{3,}", s)):
        return "填空题"
    # 5) 选择题
    if has_options and re.search(r"[（(]\s*[）)]|选择题|下列.*的是", s):
        return "单选题"
    if has_options:
        return "多选题"
    # 6) 计算题
    if re.search(r"计算|化简|解方程|求值", s):
        return "计算题"
    return "解答题"


# 客观题：机器判分，答案是文本（选项字母 / 填空词），**没有手写作答扫描**
OBJECTIVE_TYPES = ("单选题", "多选题", "判断题", "填空题")

# 平台 answer_type 编码 → 人话。语义在 2026-09-24 用 **314 道真题 / 30 场考试**
# 实测确认，分割率 100%：
#   s01Text   111 题 —— 答案全是文本、0 题有手写作答扫描  → 客观题
#   s02Image  203 题 —— 203 题**全部**有手写作答扫描      → 主观题
# 这条信号比题干启发式可靠得多，所以大类以它为准。
ANSWER_TYPE_MEANING = {
    "s01text": "客观题（机器判分，无手写作答）",
    "s02image": "主观题（人工阅卷，有手写作答）",
}


def qtype_from_answer_type(answer_type: str, stem: str, standard: str) -> str:
    """题型判定：**以题干启发式为主，平台 answer_type 只在判不出来时纠偏**。

    为什么要分主次（2026-09-24 真实数据暴露的 bug）：
    原来只用题干启发式，但很多客观题的选项是**画在图片里**的，题干里没有
    `A.` `B.` 这种文本，启发式判不出来就兜底成「解答题」——
    于是一道选择题被标成解答题。而 `core/validate.py` 的练习题闸门要求
    「生成题题型 == 原题题型」，标签错了，生成的同类题也会跟着错。

    answer_type 是平台给的，实测 100% 区分主客观（见 ANSWER_TYPE_MEANING），
    但它只说「客观 / 主观」，**不等于题型**：填空、完成句子这类题也可能
    需要手写作答（s02Image）。所以：
      · 启发式能给出具体题型（不是兜底的「解答题」）→ 尊重启发式
      · 启发式只能兜底到「解答题」时 → 用 answer_type 判断是不是客观题，
        是的话再用「选项是否在题干里」「标准答案是不是单个字母」定单选/填空
      · 编码不认识 → 保持兜底结果，不猜
    """
    guess = _guess_qtype(stem, standard or "")
    if guess != "解答题":
        return guess
    at = (answer_type or "").strip().lower()
    if at == "s01text":
        # 客观题：机器判分，答案是文本 —— 只可能是选择/填空/判断
        if re.search(r"^\s*[A-D][.．、]", _html_to_lines(stem), re.M):
            return "单选题"
        if re.fullmatch(r"[A-D]", _html_to_lines(standard).strip()):
            # 标准答案就是一个选项字母 —— 题干里的选项多半画在图里了
            return "单选题"
        return "填空题"
    return guess
