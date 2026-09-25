"""补测：覆盖第一轮验收没走到的代码路径。

第一轮 `tools/acceptance.py` 覆盖的是「归一化 → 入库 → 校验 → 统计 → 导出」这条主干。
但有几条路它没走：

  1. **通道 B 的写入链路 `sync()`** —— 411 行的适配器里最核心的函数，
     也是你给了 Cookie 之后**第一个会跑**的代码。第一轮完全没测过。
     这里用假对象（FakeStudent / FakeTopic）把它离线跑通，
     包括 40217 重试、难度未知、standard_answer 是 URL、图片落盘、去重、学科筛选。
  2. **通道 A 的 .docx 与 .pdf** —— 第一轮只测了 .txt。
  3. **Cookie 的 DPAPI 降级路径** —— keyring 可用时永远走不到那条分支。
  4. **老数据库迁移 `_migrate()`** —— 第一轮建的库本来就是新 schema。
  5. **xlsx 导出的优雅失败** —— 本机没装 openpyxl。

跑法：
    .venv/Scripts/python tools/edge_test.py

全程不联网（图片用 file:// 本地文件模拟），不碰真实库、不碰系统凭据管理器。
"""

from __future__ import annotations

import json
import os
import shutil
import sqlite3
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

TMP = Path(tempfile.mkdtemp(prefix="zx_edge_", dir=str(ROOT / "data")))
os.environ["ZX_DB_PATH"] = str(TMP / "edge.db")
os.environ["ZX_IMAGES_DIR"] = str(TMP / "images")
os.environ["ZX_TAXONOMY"] = str(ROOT / "data" / "taxonomy.yaml")
# D 段会往「降级 Cookie 文件」里写测试值再删掉。那个文件不在库里，
# 光改 ZX_DB_PATH 保护不了它 —— 必须把凭据命名空间也指到临时目录，
# 否则可能覆盖/删掉用户真实保存的 Cookie。（2026-09-24 由 mcp_e2e 踩出）
os.environ["ZX_CRED_SERVICE"] = "zhixue-wrongbook-edge"
os.environ["ZX_CRED_KEY"] = "zx_cookie_edge"
os.environ["ZX_CRED_FILE"] = str(TMP / ".session_edge.bin")

from adapters import zhixuewang as zxw                       # noqa: E402
from adapters.parsers import zhixue_export                   # noqa: E402
from core.export import wrongbook_markdown, wrongbook_xlsx   # noqa: E402
from core.store import SCHEMA, Store                         # noqa: E402

RESULTS: list[tuple[str, bool, str]] = []


def log(name: str, ok: bool, detail: str = "") -> bool:
    RESULTS.append((name, bool(ok), str(detail)[:400]))
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f"  — {detail}" if detail else ""))
    return bool(ok)


def section(t: str) -> None:
    print(f"\n{'=' * 68}\n{t}\n{'=' * 68}")


# ===========================================================================
# 假对象：模拟 zhixuewang 的返回结构（字段名照源码 models.py 抄）
# ===========================================================================
class FakeTopic:
    """模拟 ErrorBookTopic（models.py:462-480）。"""

    def __init__(self, no, **kw):
        self.dis_title_number = no
        self.content_html = kw.get("content_html", "")
        self.topic_img_url = kw.get("topic_img_url", "")
        self.image_answer = kw.get("image_answer", [])
        self.standard_answer = kw.get("standard_answer", "")
        self.answer_html = kw.get("answer_html", "")
        self.analysis_html = kw.get("analysis_html", "")
        self.difficulty = kw.get("difficulty")          # 注意：源码里是 int
        self.answer_type = kw.get("answer_type", "1")
        self.class_score_rate = kw.get("class_score_rate")
        self.score = kw.get("score")
        self.standard_score = kw.get("standard_score")
        self.paper_id = kw.get("paper_id", "")
        self.subject_name = kw.get("subject_name", "")   # 实际是试卷名
        self.topic_set_id = kw.get("topic_set_id", "")
        self.topic_analysis_img_url = kw.get("topic_analysis_img_url", "")
        self.is_correct = kw.get("is_correct", False)


class FakeSubject:
    def __init__(self, name, sid, score=100.0, code="M"):
        self.name, self.id, self.standard_score, self.code = name, sid, score, code


class FakeExam:
    def __init__(self, eid, name, create_time_ms, subjects=None,
                 grade_code="8", is_final=False):
        self.id, self.name = eid, name
        self.create_time = create_time_ms
        self.subjects = subjects or []
        self.grade_code, self.is_final = grade_code, is_final


class FakeStudent:
    """模拟 StudentAccount 里 sync() 用到的 5 个方法。

    特意做了几件事来验证降级路径：
      - E1/物理 的 paperId 调用抛 40217 → 应该自动用 topicSetId 重试成功
      - E2/数学 的调用永远抛 40217 且 E2 不是最新考试 → 没有可用兜底，
        应该记进 errors 并**继续同步下一科**（不整次中断）
    """

    def __init__(self):
        self.calls: list[tuple[str, str, str]] = []   # (exam_id, param_value, param_name)

        self.e1 = FakeExam("E1", "第一次月考", 1757865600000,
                           [FakeSubject("数学", "PAPER-M"),
                            FakeSubject("物理", "PAPER-P")])
        self.e2 = FakeExam("E2", "上学期期末", 1750291200000,
                           [FakeSubject("数学", "PAPER-M2")])
        self.exams = [self.e1, self.e2]

        # 最新考试（get_latest_exam）里的 Subject.id 来自 topicSetId（student.py:248）
        self.latest = FakeExam("E1", "第一次月考", 1757865600000,
                               [FakeSubject("数学", "TOPICSET-M"),
                                FakeSubject("物理", "TOPICSET-P")])

        self._topics = {
            ("E1", "PAPER-M"): [
                # 普通题：题干+答案+解析+学生作答+得分+难度
                FakeTopic("1", content_html="<p>已知 x^2 - 3x + m = 0 有两个不相等的实数根。</p>",
                          answer_html="m < 9/4", standard_answer="m < 9/4",
                          analysis_html="<p>由 Δ = 9 - 4m > 0 得。</p>",
                          image_answer=[], difficulty=3, class_score_rate=0.41,
                          score=3, standard_score=8),
                # standard_answer 是 URL（源码注释说的那种情况）→ 应回退到 answer_html
                FakeTopic("2", content_html="<p>求二次函数 y = x^2 - 4x + 3 的最小值。</p>",
                          standard_answer="https://www.zhixue.com/answer/abc123",
                          answer_html="最小值为 -1",
                          image_answer=["file:///" + str(TMP / "fake_ans.jpg").replace("\\", "/")],
                          difficulty=4, class_score_rate=0.33, score=0, standard_score=6),
                # 缺题干 → 解析置信度应明显偏低
                FakeTopic("3", content_html="",
                          standard_answer="", answer_html="",
                          image_answer=[], difficulty=2, score=0, standard_score=4),
            ],
            ("E1", "TOPICSET-P"): [
                FakeTopic("1", content_html="<p>R1=10Ω 与 R2=20Ω 串联，求电流。</p>",
                          answer_html="I = 0.2A", standard_answer="I = 0.2A",
                          difficulty=3, class_score_rate=0.82, score=0, standard_score=6),
                FakeTopic("2", content_html="<p>求滑轮组机械效率。</p>",
                          answer_html="80%", standard_answer="80%",
                          difficulty=4, class_score_rate=0.55, score=2, standard_score=6),
            ],
        }

    # --- sync() 会用到的接口 ---
    def get_exams(self):
        return list(self.exams)

    def get_exam(self, exam_id):
        for e in self.exams:
            if e.id == exam_id:
                return e
        return None

    def get_subjects(self, exam):
        return list(exam.subjects)

    def get_latest_exam(self):
        return self.latest

    def get_errorbook(self, exam_id, param):
        self.calls.append((exam_id, param,
                           "topicSetId" if param.startswith("TOPICSET") else "paperId"))
        if (exam_id, param) in self._topics:
            return list(self._topics[(exam_id, param)])
        # 模拟库的行为：errorCode != 0 时抛裸 Exception，data 是 dict（student.py:831）
        raise Exception({"errorCode": 40217,
                         "errorInfo": "暂时未收集到试题信息,无法查看", "result": ""})


# ===========================================================================
def main() -> int:
    # 造一个假的学生作答图片，供 file:// 下载测试
    (TMP / "fake_ans.jpg").write_bytes(b"\xff\xd8\xff\xe0FAKEJPEG")

    student = FakeStudent()

    # ------------------------------------------------------------------ A
    section("A. 通道 B 的写入链路 sync()（第一轮完全没覆盖）")
    zxw.login = lambda raw: student            # 猴补丁：跳过真实登录
    zxw.to_student = lambda acc: acc

    store = Store(TMP / "edge.db", TMP / "images")
    report = zxw.sync(store, "token=x; loginUserName=u; userName=u",
                      max_exams=10, download=True)

    check_a1 = log("sync 完成且不抛异常", isinstance(report, dict),
                   f"扫描考试 {report['exams_scanned']} 场 / "
                   f"学科 {report['subjects_scanned']} 个")
    log("扫描到 2 场考试、3 个学科条目", report["exams_scanned"] == 2
        and report["subjects_scanned"] == 3, "")
    log("入库题数正确（E1 数学 3 题 + E1 物理 2 题 = 5）",
        report["added"] == 5 and store.counts()["total"] == 5,
        f"added={report['added']}，库内总数={store.counts()['total']}")

    # 40217 重试
    retried = [c for c in student.calls if c[2] == "topicSetId"]
    log("paperId 失败后自动用 topicSetId 重试并成功", len(retried) == 1,
        f"重试调用={retried}")
    phys = [e for e in report["per_subject"]
            if e["exam"] == "第一次月考" and e["subject"] == "物理"]
    log("重试生效被写进报告（param_used）",
        phys and phys[0].get("param_used") == "topicSetId(重试)",
        f"param_used={phys[0].get('param_used') if phys else None}")
    log("重试原因也记进了 notes",
        any("param_note" in n for n in report["notes"]), "")

    # 一个学科失败不该拖垮整次同步
    e2 = [e for e in report["per_subject"] if e["exam"] == "上学期期末"]
    log("E2/数学 无兜底时记错误但不中断后续同步",
        e2 and e2[0]["error"] and "40217" in e2[0]["error"],
        f"error={e2[0]['error'] if e2 else None}")
    log("E2 用的是它自己的 paperId，没有借用别的考试的 topicSetId",
        not any(c[0] == "E2" and c[2] == "topicSetId" for c in student.calls),
        f"E2 的调用={[c for c in student.calls if c[0] == 'E2']}")

    # 难度刻度：三种 scale 的行为 + 报告自洽 + 未核实那条路也要真跑到
    from core.config import Config as _Cfg

    c_unk = _Cfg({"difficulty": {"scale": "unknown"}})
    c_raw = _Cfg({"difficulty": {"scale": "raw", "assumed_max": 3}})
    c_01 = _Cfg({"difficulty": {"scale": "0-1"}})
    log("scale=raw + assumed_max=3：难度 3 → 1.0，难度 0 → 0.0",
        c_raw.normalize_difficulty(3) == (1.0, "0-1")
        and c_raw.normalize_difficulty(0) == (0.0, "0-1"),
        f"3→{c_raw.normalize_difficulty(3)} 0→{c_raw.normalize_difficulty(0)}")
    log("scale=unknown：返回原值 + 'unknown'，不假装能比",
        c_unk.normalize_difficulty(3) == (3.0, "unknown"),
        f"{c_unk.normalize_difficulty(3)}")
    log("scale=0-1：按原样夹到 0~1",
        c_01.normalize_difficulty(0.42) == (0.42, "0-1"), "")
    log("难度为 None / 非数字时安全返回 (None, 'unknown')",
        c_raw.normalize_difficulty(None) == (None, "unknown")
        and c_raw.normalize_difficulty("abc") == (None, "unknown"), "")
    log("difficulty_is_verified 随配置正确翻转",
        c_raw.difficulty_is_verified and not c_unk.difficulty_is_verified
        and c_01.difficulty_is_verified,
        f"raw={c_raw.difficulty_is_verified} unknown={c_unk.difficulty_is_verified} "
        f"0-1={c_01.difficulty_is_verified}")

    # 报告必须与配置自洽：已核实就不该再报「未核实」，反之必须报
    verified = report["difficulty_scale_verified"]
    warned = any("难度刻度未核实" in w for w in report.get("warnings", []))
    log("报告里的难度刻度状态与配置自洽（未核实 → 必须给警告）",
        warned == (not verified),
        f"scale={report['difficulty_scale']} verified={verified} warned={warned}")

    # 未核实那条路也要真跑到 —— 临时把 config 换成 scale=unknown 再同步一次
    _real_get_config = zxw.get_config
    _base = _real_get_config()
    zxw.get_config = lambda *a, **k: _Cfg({
        "difficulty": {"scale": "unknown"},
        "data": _base._d["data"], "sync": _base._d["sync"]})
    try:
        store_u = Store(TMP / "unverified.db", TMP / "images_unver")
        rep_u = zxw.sync(store_u, "token=x; loginUserName=u; userName=u",
                         max_exams=10, download=False)
        log("刻度未核实时：报告明确给出警告（不假装通过）",
            any("难度刻度未核实" in w for w in rep_u.get("warnings", [])),
            f"warnings={rep_u.get('warnings')}")
        log("刻度未核实时：每题的原始难度值被记进 notes（供核实用）",
            any("difficulty_scale_unknown" in n for n in rep_u["notes"]),
            f"例子={[n.get('difficulty_scale_unknown') for n in rep_u['notes'] if 'difficulty_scale_unknown' in n][:3]}")
        store_u.close()
    finally:
        zxw.get_config = _real_get_config

    # standard_answer 是 URL
    log("standard_answer 是 URL 时回退到 answer_html",
        any("standard_answer_is_url" in n for n in report["notes"]), "")
    q2 = [q for q in store.query() if "二次函数" in q.stem_text]
    log("该题的标准答案正文取到了 answer_html 而不是链接",
        q2 and q2[0].answer.standard == "最小值为 -1",
        f"standard={q2[0].answer.standard!r}" if q2 else "没找到题")

    # 图片落盘
    imgs = list((TMP / "images").glob("*.jpg"))
    log("file:// 图片被复制进 images 目录（不是只记路径）", len(imgs) >= 1,
        f"{[i.name for i in imgs]}")
    q2b = store.get(q2[0].fingerprint()) if q2 else None
    log("学生作答图片记到了该题上", q2b and len(q2b.answer.student_images) == 1,
        f"{q2b.answer.student_images if q2b else None}")

    # 解析置信度分级
    low = [q for q in store.query(exclude_low_confidence=False)
           if q.parse_confidence < 0.5]
    log("缺题干的题解析置信度偏低", len(low) == 1,
        f"{len(low)} 道低置信度，值={[q.parse_confidence for q in low]}")

    # 二次同步 = 去重 + 增量
    report2 = zxw.sync(store, "token=x; loginUserName=u; userName=u",
                       max_exams=10, download=True)
    log("二次同步不新增（指纹去重）",
        report2["added"] == 0 and report2["updated"] == 5
        and store.counts()["total"] == 5,
        f"added={report2['added']} updated={report2['updated']} "
        f"总数={store.counts()['total']}")

    # 学科筛选
    store2 = Store(TMP / "edge2.db", TMP / "images2")
    r3 = zxw.sync(store2, "token=x; loginUserName=u; userName=u",
                  subjects=["数学"], max_exams=10, download=False)
    log("学科筛选生效（只同步数学）",
        r3["added"] == 3 and r3["subjects_scanned"] == 2,
        f"added={r3['added']}，扫描学科={r3['subjects_scanned']}（E1数学+E2数学）")
    store2.close()

    # 同步日志
    logs = store.sync_logs()
    log("同步日志按学科逐条记录", len(logs) >= 3, f"{len(logs)} 条")

    # ------------------------------------------------------------------ B
    section("B. 通道 A：.docx 解析（第一轮只测了 .txt）")
    try:
        import docx
        d = docx.Document()
        for line in [
            "第一次月考 数学",
            "1. 已知关于 x 的一元二次方程 x^2 - 3x + m = 0 有两个不相等的实数根。求 m 的取值范围。",
            "答案：m < 9/4",
            "解析：由判别式 Δ = 9 - 4m > 0 得 m < 9/4。",
            "学生答案：m <= 9/4",
            "得分：3/8",
            "难度：0.62",
            "2. 计算 (-2)^3 + 5 的值",
            "答案：-3",
            "学生答案：-3",
            "得分：5/5",
        ]:
            d.add_paragraph(line)
        docx_path = TMP / "export.docx"
        d.save(str(docx_path))

        qs, stats = zhixue_export.parse_file(docx_path, subject="数学",
                                            exam_name="第一次月考",
                                            exam_date="2026-09-15")
        log("docx 解析出 2 道题", len(qs) == 2,
            f"{len(qs)} 题，方式={stats['extract_method']}")
        log("docx 的标准答案/解析/学生作答/得分/难度全部解析到位",
            qs[0].answer.standard == "m < 9/4"
            and qs[0].answer.analysis_html
            and qs[0].answer.student_text == "m <= 9/4"
            and qs[0].score.got == 3 and qs[0].question.difficulty == 0.62,
            f"答案={qs[0].answer.standard!r} 得分={qs[0].score.got}/{qs[0].score.full} "
            f"难度={qs[0].question.difficulty}")
        log("字段齐备的题解析置信度为 1.0", qs[0].parse_confidence == 1.0,
            f"{qs[0].parse_confidence}")
        log("题型按题干启发式判断（第二题含「计算」→ 计算题）",
            qs[1].question.type == "计算题", f"{qs[1].question.type}")
        log("docx 的难度刻度标为 0-1（值域在 0~1 内）",
            qs[0].question.difficulty_scale == "0-1",
            f"{qs[0].question.difficulty_scale}")
    except ImportError as exc:
        log("python-docx 可用", False, str(exc))

    # ------------------------------------------------------------------ C
    section("C. 通道 A：.pdf 解析（自建最小 PDF，验证 pypdf 通链路）")
    pdf_path = TMP / "export_ascii.pdf"
    pdf_path.write_bytes(_build_minimal_pdf([
        "1. Solve the equation x^2 - 4x + 3 = 0 .",
        "Answer: x = 1 or x = 3",
        "Score: 3/8",
        "2. Compute (-2)^3 + 5 .",
        "Answer: -3",
        "Score: 5/5",
    ]))
    try:
        qs, stats = zhixue_export.parse_file(pdf_path, subject="数学",
                                            exam_name="ASCII 导出测试",
                                            exam_date="2026-09-15")
        log("PDF 文本层被成功抽取并切出题块", len(qs) == 2,
            f"{len(qs)} 题，方式={stats['extract_method']}")
        log("没有中文标签时答案不会凭空产生（standard 为 None）",
            qs[0].answer.standard is None,
            f"standard={qs[0].answer.standard!r}")
        log("但得分仍能按数字规则解析出来", qs[0].score.got == 3,
            f"{qs[0].score.got}/{qs[0].score.full}")
        log("缺少答案/作答/难度 → 置信度 0.40，会被统计自动排除",
            qs[0].parse_confidence == 0.40 and stats["low_confidence"] == 2,
            f"置信度={qs[0].parse_confidence}，低置信度题数={stats['low_confidence']}")
    except Exception as exc:
        log("PDF 解析", False, f"{type(exc).__name__}: {exc}")

    # 扫描版 PDF（无文本层）必须明确报错，不能假装成功
    blank_pdf = TMP / "scanned_like.pdf"
    blank_pdf.write_bytes(_build_minimal_pdf([], draw_only=True))
    try:
        zhixue_export.parse_file(blank_pdf, subject="数学", exam_name="扫描版")
        log("无文本层 PDF 应明确报错", False, "居然没报错")
    except RuntimeError as exc:
        log("无文本层 PDF 给出明确错误（不伪造结果）",
            "未提取到任何文本" in str(exc), str(exc)[:120])
    except Exception as exc:
        log("无文本层 PDF 给出明确错误", False, f"抛了别的异常：{type(exc).__name__}")

    # ------------------------------------------------------------------ D
    section("D. Cookie 存储：DPAPI 降级路径 + 输入校验（不碰系统凭据管理器）")
    from adapters import session as sess

    # 先确认隔离生效 —— 没有这条，下面的写/删就是在动用户真实凭据
    log("凭据命名空间已隔离到临时目录（不会碰真实 Cookie）",
        sess.SERVICE_NAME != "zhixue-wrongbook"
        and sess.KEY_NAME != "zx_cookie"
        and str(TMP) in str(sess.FALLBACK_PATH),
        f"service={sess.SERVICE_NAME} file={sess.FALLBACK_PATH}")

    try:
        sess._file_set("token=abc==; loginUserName=u; userName=u")
        got = sess._file_get()
        log("DPAPI 加密往返一致（值里的 '=' 也没丢）",
            got == "token=abc==; loginUserName=u; userName=u", f"{got!r}")
        raw = sess.FALLBACK_PATH.read_bytes()
        log("落盘内容不是明文（是 base64 后的密文）",
            b"loginUserName" not in raw, f"前 24 字节={raw[:24]!r}")
        sess._file_delete()
        log("删除后读不到", sess._file_get() is None, "")
    except Exception as exc:
        log("DPAPI 往返", False, f"{type(exc).__name__}: {exc}")

    for bad, why in [("", "空字符串"), ("没有等号的一串东西", "格式不对")]:
        try:
            sess.set_cookie(bad)
            log(f"set_cookie 拒绝{why}", False, "居然接受了")
        except (ValueError, RuntimeError) as exc:
            log(f"set_cookie 拒绝{why}", True, str(exc)[:80])

    # ------------------------------------------------------------------ E
    section("E. 老数据库迁移 _migrate()（第一轮建的库本来就是新 schema）")
    old_db = TMP / "old_schema.db"
    old_schema = "\n".join(l for l in SCHEMA.splitlines()
                           if "difficulty_scale" not in l)
    conn = sqlite3.connect(str(old_db))
    conn.executescript(old_schema)
    conn.execute(
        "INSERT INTO questions (id, fingerprint, source, source_version, "
        "fetched_at, first_seen_at, parse_confidence, subject, exam_name) "
        "VALUES ('old-1','fp-old','api','v0','2026-01-01T00:00:00',"
        "'2026-01-01T00:00:00',0.9,'数学','旧考试')")
    conn.commit()
    cols_before = {r[1] for r in conn.execute("PRAGMA table_info(questions)")}
    conn.close()
    log("迁移前：老库确实没有 difficulty_scale 列",
        "difficulty_scale" not in cols_before, "")

    s_old = Store(old_db, TMP / "images_old")
    cols_after = {r["name"] for r in
                  s_old.conn.execute("PRAGMA table_info(questions)")}
    log("迁移后：自动补上新列",
        "difficulty_scale" in cols_after and "first_seen_at" in cols_after,
        f"新增={sorted(cols_after - cols_before)}")
    log("迁移不丢老数据", s_old.counts()["total"] == 1,
        f"总数={s_old.counts()['total']}")
    q_old = s_old.query(exclude_low_confidence=False)
    log("老行的新列有合理默认值", q_old and q_old[0].question.difficulty_scale == "unknown",
        f"difficulty_scale={q_old[0].question.difficulty_scale if q_old else None}")
    s_old.close()

    # ------------------------------------------------------------------ F
    section("F. 导出：xlsx（装了/没装 openpyxl 两条路）+ 错题本 Markdown")
    # 这一段原来写死了「本机没装 openpyxl，所以应该报错」。
    # 2026-09-24 把 openpyxl 装上之后，那条断言直接变成假失败 ——
    # 测试不该假设运行环境。改成按实际有没有装分流验两条路。
    try:
        import openpyxl
        has_openpyxl = True
    except ImportError:
        has_openpyxl = False

    if has_openpyxl:
        p = Path(wrongbook_xlsx(store.query(), TMP / "wb.xlsx"))
        ok_size = p.exists() and p.stat().st_size > 0
        # 不只是「文件写出来了」，还要能读回来、表头和数据行对得上
        wb = openpyxl.load_workbook(str(p))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        log("装了 openpyxl → xlsx 真写出且能读回",
            ok_size and len(rows) >= 2 and rows[0][0] == "序号",
            f"{p.stat().st_size} B，{len(rows)} 行（含表头），表头首列={rows[0][0]!r}")
    else:
        try:
            wrongbook_xlsx(store.query(), TMP / "wb.xlsx")
            log("未装 openpyxl 时应明确报错", False, "居然成功了")
        except RuntimeError as exc:
            log("未装 openpyxl 时给出人话提示而不是崩溃",
                "openpyxl" in str(exc) and "pip install" in str(exc), str(exc)[:110])
        except Exception as exc:
            log("未装 openpyxl 时给出人话提示", False,
                f"抛了别的异常：{type(exc).__name__}: {exc}")

    md = wrongbook_markdown(store.query(), title="补测错题本")
    (TMP / "wb.md").write_text(md, encoding="utf-8")
    log("错题本 Markdown 含来源通道列与低置信度说明",
        "来源" in md and "parse_confidence" in md, f"{len(md)} 字")

    # ------------------------------------------------------------------ G
    section("G. 跨考试 id 唯一性（回归：真实同步时整批入库失败）")
    # 回归背景（2026-09-24，真实账号同步）：
    #   第一场「20260918八年级周测」数学 15 题全部入库；
    #   第二场「20260911八年级周测」数学 12 题**全部**报
    #   `UNIQUE constraint failed: questions.id`。
    # 根因：make_id 是 zx_<年>_<学科>_<序号>，而 seq 是「每场考试每学科」
    #   内部从 1 重来的 —— 同年两场数学考试的第 1 题算出了同一个 id。
    from adapters.parsers.normalize import make_id

    def _ms(y, m, d):
        return int(datetime(y, m, d, tzinfo=timezone.utc).timestamp() * 1000)

    id_a = make_id("数学", "20260918八年级周测", "2026-09-18", 1)
    id_b = make_id("数学", "20260911八年级周测", "2026-09-11", 1)
    log("同年、同学科、不同考试的第 1 题不再算出同一个 id",
        id_a != id_b, f"{id_a} vs {id_b}")
    log("同一场考试内不同序号仍然不同",
        make_id("数学", "X", "2026-09-18", 1) != make_id("数学", "X", "2026-09-18", 2), "")
    log("id 生成稳定可复现（重复同步不会换来换去）",
        id_a == make_id("数学", "20260918八年级周测", "2026-09-18", 1), "")
    log("考试名/日期缺失时也能生成合法 id",
        bool(make_id("数学", "", None, 1).startswith("zx_0000_math_")),
        make_id("数学", "", None, 1))

    class _TwoExamStudent:
        """两场考试、同学科、同年，且都成功返回题目 —— 精确复现撞车场景。"""

        def __init__(self):
            self.e1 = FakeExam("G1", "20260918八年级周测", _ms(2026, 9, 18),
                               [FakeSubject("数学", "PAPER-G1")])
            self.e2 = FakeExam("G2", "20260911八年级周测", _ms(2026, 9, 11),
                               [FakeSubject("数学", "PAPER-G2")])
            self.exams = [self.e1, self.e2]
            self.latest = self.e1

        def get_exams(self):
            return list(self.exams)

        def get_exam(self, exam_id):
            return next((e for e in self.exams if e.id == exam_id), None)

        def get_subjects(self, exam):
            return list(exam.subjects)

        def get_latest_exam(self):
            return self.latest

        def get_errorbook(self, exam_id, param):
            n = 1 if exam_id == "G1" else 2      # 两场都有题，且序号都从 1 开始
            return [FakeTopic(str(i),
                              content_html=f"<p>{exam_id} 第 {i} 题：求 x 的值。</p>",
                              answer_html=f"x={i}", standard_answer=f"x={i}",
                              difficulty=3, score=0, standard_score=5)
                    for i in range(1, n + 1)]

    store_g = Store(TMP / "idcollide.db", TMP / "images_idc")
    zxw.login = lambda raw: _TwoExamStudent()
    zxw.to_student = lambda acc: acc
    rep_g = zxw.sync(store_g, "token=x; loginUserName=u; userName=u",
                     max_exams=10, download=False)

    log("两场考试的同学科题目全部入库（不再整批失败）",
        rep_g["added"] == 3 and rep_g["failed"] == 0,
        f"added={rep_g['added']} failed={rep_g['failed']} "
        f"库内总数={store_g.counts()['total']}")
    log("没有任何一道题进 errors",
        not [e for e in rep_g["errors"] if "UNIQUE" in str(e.get("error", ""))],
        f"errors={rep_g['errors'][:2]}")
    log("库内 id 全部唯一",
        len({q.id for q in store_g.query(exclude_low_confidence=False)})
        == store_g.counts()["total"], "")
    log("没有触发 id 兜底改名 —— 说明根因已修，不是被兜底掩盖",
        store_g.id_fallbacks == [], f"{store_g.id_fallbacks}")
    store_g.close()

    # 兜底路径本身也要能用：手工塞一条占住目标 id，再插同 id 的新题
    store_g2 = Store(TMP / "idcollide2.db", TMP / "images_idc2")
    conn = store_g2.conn
    conn.execute(
        "INSERT INTO questions (id, fingerprint, source, source_version, "
        "fetched_at, first_seen_at, parse_confidence, subject, exam_name) "
        "VALUES ('taken-1','fp-taken','manual','','2026-01-01T00:00:00',"
        "'2026-01-01T00:00:00',1.0,'数学','占位考试')")
    conn.commit()
    from core.models import (AnswerPart, Exam, QuestionPart, Score,
                             WrongQuestion)
    q_dup = WrongQuestion(
        id="taken-1", source="manual", subject="数学",
        exam=Exam(name="另一场考试", date="2026-03-01"),
        question=QuestionPart(stem_html="<p>占位 id 的另一道题</p>"),
        answer=AnswerPart(), score=Score(),
    )
    try:
        res = store_g2.upsert(q_dup)
        log("id 被占用时自动换一个唯一 id 而不是抛异常",
            res == "added" and store_g2.id_fallbacks
            and store_g2.counts()["total"] == 2,
            f"result={res} fallbacks={store_g2.id_fallbacks} "
            f"总数={store_g2.counts()['total']}")
        new_id = store_g2.query(exclude_low_confidence=False)
        used = [q.id for q in new_id if q.stem_text.startswith("占位 id")]
        log("换出来的 id 带指纹后缀、可追溯",
            used and used[0].startswith("taken-1-"), f"{used}")
    except Exception as exc:
        log("id 被占用时自动换一个唯一 id", False,
            f"抛异常 {type(exc).__name__}: {exc}")
    store_g2.close()

    # ------------------------------------------------------------------ H
    section("H. 学生作答图片：JSON 字符串形态（回归：作答图一张都存不下来）")
    # 回归背景（2026-09-24，真实账号）：
    #   库把接口的 imageAnswer 原样透传（student.py:844），而它实际是**JSON 编码
    #   的字符串**：'["https://...jpg?Expires=...&Signature=..."]'，
    #   但类型标注写的是 List[str]。
    #   旧代码「是 str 就包成单元素列表」，得到 ['["https://..."]']，
    #   带着方括号和引号去 requests.get 必然失败 ——
    #   **学生手写作答图片一张都没存下来**，而它正是错因分析最关键的证据。
    from adapters.zhixuewang import _as_url_list

    log("JSON 数组字符串被正确拆成 URL 列表",
        _as_url_list('["https://a.jpg"]') == ["https://a.jpg"], "")
    log("多个元素的 JSON 数组",
        _as_url_list('["https://a.jpg","https://b.jpg"]')
        == ["https://a.jpg", "https://b.jpg"], "")
    log("裸 URL 仍然可用（不破坏原有行为）",
        _as_url_list("https://a.jpg") == ["https://a.jpg"], "")
    log("带签名参数的 URL 不被破坏",
        _as_url_list('["https://a.jpg?Expires=1&Signature=x%2Fy"]')
        == ["https://a.jpg?Expires=1&Signature=x%2Fy"], "")
    log("空值 / None / 空数组都安全",
        _as_url_list(None) == [] and _as_url_list("") == []
        and _as_url_list([]) == [] and _as_url_list("[]") == [], "")
    log("坏 JSON 不抛异常，退回原文",
        _as_url_list("[坏JSON") == ["[坏JSON"], "")
    log("列表里套 JSON 串也能剥开",
        _as_url_list(['["https://a.jpg"]']) == ["https://a.jpg"], "")

    class _JsonAnsStudent:
        """image_answer 是 JSON 字符串 —— 精确复现真实形态。"""

        def __init__(self):
            self.e1 = FakeExam("H1", "JSON作答图回归", _ms(2026, 8, 1),
                               [FakeSubject("数学", "PAPER-H1")])
            self.exams = [self.e1]
            self.latest = self.e1
            url = "file:///" + str(TMP / "fake_ans.jpg").replace("\\", "/")
            self._json_ans = json.dumps([url], ensure_ascii=False)

        def get_exams(self):
            return list(self.exams)

        def get_exam(self, exam_id):
            return next((e for e in self.exams if e.id == exam_id), None)

        def get_subjects(self, exam):
            return list(exam.subjects)

        def get_latest_exam(self):
            return self.latest

        def get_errorbook(self, exam_id, param):
            return [FakeTopic("1", content_html="<p>JSON 作答图回归题</p>",
                              answer_html="x=1", standard_answer="x=1",
                              image_answer=self._json_ans,
                              difficulty=3, score=0, standard_score=5)]

    store_h = Store(TMP / "jsonans.db", TMP / "images_jsonans")
    zxw.login = lambda raw: _JsonAnsStudent()
    zxw.to_student = lambda acc: acc
    rep_h = zxw.sync(store_h, "token=x; loginUserName=u; userName=u",
                     max_exams=10, download=True)
    got = [q for q in store_h.query(exclude_low_confidence=False)
           if "JSON 作答图" in q.stem_text]
    # 注意：图片落到哪由 config 的 images_dir 决定（模块顶部设的 ZX_IMAGES_DIR），
    # 不是 Store(images_dir=...) 那个参数 —— 别搞混。
    from core.config import get_config as _get_config
    img_root = _get_config().path("images_dir")
    log("JSON 字符串形态的作答图能被下载并落到本地",
        bool(got) and len(got[0].answer.student_images) == 1
        and (img_root / got[0].answer.student_images[0]).exists(),
        f"{got[0].answer.student_images if got else None} @ {img_root}")
    log("作答图记的是本地文件名，而不是 URL / JSON 原文",
        bool(got) and bool(got[0].answer.student_images)
        and not got[0].answer.student_images[0].startswith(("http", "[")),
        f"{got[0].answer.student_images if got else None}")
    log("没有 image_failed / student_images_download_failed 记录",
        not [n for n in rep_h["notes"]
             if "image_failed" in n or "student_images_download_failed" in n], "")
    store_h.close()

    # ------------------------------------------------------------------ I
    section("I. 指纹稳定性（回归：改题型规则导致 11 道老题被重复入库）")
    # 回归背景（2026-09-24）：fingerprint 原来把 `question.type` 也算进去，
    # 而题型是**推导出来的标签**（API 通道由 answer_type+启发式算，
    # 导出通道由文件文本算）。于是题型判定规则一改，11 道已入库的题指纹全变、
    # 被当成新题重复插入，题库凭空多出 11 条 ——
    # 正好违背 fingerprint() 自己承诺的「跨通道同一道题应合并成一条」。
    from core.models import AnswerPart as _AP
    from core.models import Exam as _Ex
    from core.models import QuestionPart as _QP
    from core.models import Score as _Sc
    from core.models import WrongQuestion as _WQ

    def _mk(qtype, stem="<p>同一道题的题干</p>", exam="同一场考试"):
        return _WQ(id="x", source="manual", subject="数学",
                   exam=_Ex(name=exam, date="2026-01-01"),
                   question=_QP(stem_html=stem, type=qtype),
                   answer=_AP(standard="42"), score=_Sc())

    log("指纹不含题型（改判定规则不会让老题变成新题）",
        _mk("解答题").fingerprint() == _mk("单选题").fingerprint(), "")
    log("指纹对题干敏感（内容变了就该是新题）",
        _mk("解答题").fingerprint() != _mk("解答题", "<p>换了题干</p>").fingerprint(), "")
    log("指纹对考试敏感（不同考试的同题干题不合并）",
        _mk("解答题").fingerprint() != _mk("解答题", exam="另一场考试").fingerprint(), "")
    log("指纹对标准答案敏感",
        _mk("解答题").fingerprint() != _WQ(
            id="x", source="manual", subject="数学",
            exam=_Ex(name="同一场考试", date="2026-01-01"),
            question=_QP(stem_html="<p>同一道题的题干</p>"),
            answer=_AP(standard="43"), score=_Sc()).fingerprint(), "")

    # 集成：同一批数据连同步两次，第二次不该新增
    store_f = Store(TMP / "fpstable.db", TMP / "images_fp")
    zxw.login = lambda raw: FakeStudent()
    zxw.to_student = lambda acc: acc
    r1 = zxw.sync(store_f, "token=x; loginUserName=u; userName=u",
                  max_exams=10, download=False)
    r2 = zxw.sync(store_f, "token=x; loginUserName=u; userName=u",
                  max_exams=10, download=False)
    log("二次同步零新增（指纹跨次稳定）",
        r1["added"] == 5 and r2["added"] == 0 and r2["updated"] == 5,
        f"第一次 added={r1['added']}；第二次 added={r2['added']} "
        f"updated={r2['updated']} 总数={store_f.counts()['total']}")
    store_f.close()

    # ------------------------------------------------------------------ J
    section("J. 公式还原 + 指纹过期自检（回归：公式全丢 / 逻辑一改就重复入库）")
    # 回归背景（2026-09-24，真实数学题）：
    #   平台把公式渲染成 <img>，但**把 LaTeX 原文放在 data-latex 属性里**。
    #   html_to_text 直接剥标签，公式全丢 —— 一道题的真实题干
    #   「在▱ABCD中，AB=4，BC=6，∠B=60°，…求AE+½CF最大值」
    #   退化成「如图，在中，，，，点、分别在、上」，
    #   标准答案从 √31+√7 退化成单个 "+"，解析里 96 段 LaTeX 全部被丢掉。
    from core.models import html_to_text as _h2t

    log("带 data-latex 的公式图被还原成 LaTeX 原文",
        _h2t('<p>在<img data-latex="%20%5Cparallelogram%20"/>ABCD中</p>')
        == "在 \\parallelogram ABCD中", "")
    log("普通图片仍然被丢掉（不破坏原有行为）",
        _h2t('<p>文字</p><img src="a.png"/>') == "文字", "")
    log("分式等复杂 LaTeX 能还原",
        _h2t('<img data-latex="%5Cfrac%7B1%7D%7B2%7D"/>') == "\\frac{1}{2}", "")
    log("没有 data-latex 属性的 img 不会产生垃圾字符",
        _h2t('<img src="x.png" type="ocr"/>') == "", "")
    log("HTML 实体仍正常解码",
        _h2t("<p>a &lt; b &amp; c</p>") == "a < b & c", "")

    # 指纹过期自检：库里存的是旧逻辑算的指纹时，必须能检测出来
    store_j = Store(TMP / "stale.db", TMP / "images_stale")
    store_j.conn.execute(
        "INSERT INTO questions (id, fingerprint, source, source_version, "
        "fetched_at, first_seen_at, parse_confidence, subject, exam_name, "
        "stem_html) VALUES ('s1','旧逻辑算出来的指纹','api','','2026-01-01T00:00:00',"
        "'2026-01-01T00:00:00',1.0,'数学','某考试','<p>题干</p>')")
    store_j.conn.commit()
    log("能检测出「指纹与当前逻辑不一致」的老行",
        store_j.stale_fingerprint_count() == 1,
        f"stale={store_j.stale_fingerprint_count()}")
    store_j.close()

    # 一致的库不该误报
    store_j2 = Store(TMP / "stale2.db", TMP / "images_stale2")
    zxw.login = lambda raw: FakeStudent()
    zxw.to_student = lambda acc: acc
    zxw.sync(store_j2, "token=x; loginUserName=u; userName=u",
             max_exams=10, download=False)
    log("刚同步过的库不误报（stale=0）",
        store_j2.stale_fingerprint_count() == 0,
        f"stale={store_j2.stale_fingerprint_count()} / 总数 {store_j2.counts()['total']}")
    rep_j = zxw.sync(store_j2, "token=x; loginUserName=u; userName=u",
                     max_exams=10, download=False)
    log("指纹一致时同步报告不出现「重复入库」警告",
        not any("重复入库" in w for w in rep_j.get("warnings", [])),
        f"warnings={rep_j.get('warnings')}")
    store_j2.close()

    # ---------------------------------------------------------- 题型判定
    section("K. 题型判定（回归：探究题带选项被误判 / 填空下划线丢失）")
    gq = zhixue_export._guess_qtype

    # 缺陷 1（2026-09-24 真实数据）：探究题的小问自带 A/B/C 选项，
    # 原来的「先看选项」把整题判成选择题。实测 4 道中招。
    log("探究题带选项 → 实验探究题（曾被判成单选题）",
        gq('<p>利用如图甲所示的实验装置观察水的沸腾：</p>'
           '<p>(1)在安装器材时出现如图甲所示的情形，应调 <u class="blank b1"> </u></p>'
           '<p>(2)① <u class="blank b2"> </u> ②C；</p>'
           '<p>(3)A．甲 B．乙 C．丙</p>', '') == "实验探究题")
    log("多选题式探究题 → 实验探究题（曾被判成多选题）",
        gq('<p>用图甲所示的实验装置做“探究平面镜成像特点”实验。</p>'
           '<p>(1)采用玻璃板代替平面镜，主要目的是 <u class="blank b1"> </u></p>'
           '<p>(2)实验时，将蜡烛A放在玻璃板前，蜡烛B <u class="blank b2"> </u></p>', '')
        == "实验探究题")

    # 缺陷 2：填空下划线在 HTML 里是 <u class="blank …">，剥标签后就没了。
    log("HTML 下划线填空 → 填空题（曾被判成解答题）",
        gq('<p>其中两个像由于光射到玻璃发生折射形成的，分别是 '
           '<u class="blank blank1 _blank_1"> </u>号；另外两个像的形成原因是由于光的 '
           '<u class="blank blank2 _blank_2"> </u>.</p>', '') == "填空题")
    log("普通下划线填空 → 填空题",
        gq('<p>该函数的最大值是 __________ .</p>', '') == "填空题")

    # 护栏：有选项就不算填空，否则「完形填空」会被吃成填空题。
    log("带选项的完形填空不会被判成填空题",
        gq('<p>完形填空</p><p>A. a</p><p>B. b</p><p>C. c</p>', '') != "填空题")

    # 顺序：作图题的小问里常有填空横线，不能被填空题规则吃掉。
    log("作图题优先于填空题",
        gq('<p>按要求作图（请保留作图痕迹）</p>'
           '<p>(1)一束光与平面镜成 <u class="blank b1"> </u> 角</p>'
           '<p>(2)画出反射光线</p>', '') == "作图题")

    # 兜底：什么都没命中时仍然是「解答题」，不是瞎猜。
    log("无标记的解答题仍兜底为解答题",
        gq('<p>如图，在△ABC中，求∠A的度数。</p>', '') == "解答题")

    # ------------------------------------------------------ 答案比对判定
    section("L. 答案比对（回归：50 字证明题被硬判成 mismatch 并写进练习历史）")
    from core.validate import compare_solution as cmp_sol

    proof_sub = ("证明：原式=(4n^2+4n+1)-(4n^2-4n+1)=8n。因为n为整数，所以8n是8的倍数，"
                 "即(2n+1)^2-(2n-1)^2的值一定是8的倍数。")
    proof_std = ("证明：原式=(4n^2+4n+1)-(4n^2-4n+1)=8n，因为n为整数，所以8n是8的倍数。")
    log("成段文字的证明题判为 undecidable（曾被硬判 mismatch）",
        cmp_sol(proof_sub, proof_std)["verdict"] == "undecidable",
        cmp_sol(proof_sub, proof_std)["verdict"])
    log("英文开放题长答案判为 undecidable",
        cmp_sol("It can help us develop a healthy attitude towards labour and master skills.",
                "It can help us develop a healthy attitude towards labour.")["verdict"]
        == "undecidable")
    log("短答案仍硬判：一致 → match",
        cmp_sol("折射;反射", "折射;反射")["verdict"] == "match")
    log("短答案仍硬判：不一致 → mismatch",
        cmp_sol("③④;直线传播", "②③;反射")["verdict"] == "mismatch")
    log("选项字母不一致仍判 mismatch",
        cmp_sol("A", "B")["verdict"] == "mismatch")
    log("短中文填空答案（6 字以内）不被当成成段文字",
        cmp_sol("近视;25;短", "近视;25;短")["verdict"] == "match")

    # 缺陷 3（2026-09-24 修）：numeric_match 原来排在成段文字判定**之前**。
    # 一道证明题，宿主写「…=8n，n为整数，故8n是8的倍数」、标准答案写
    # 「…=8n，因为n为整数，所以8n是8的倍数」—— 两边数字序列恰好都是
    # [4,2,4,1,4,2,-4,1,8,8,8]，于是被判 numeric_match。
    # 这跟本函数文档写的「任一侧是成段文字就不硬判对错」直接矛盾：
    # 数字序列撞车纯属偶然，撞上了也不代表推导对。
    v = cmp_sol("证明：原式=(4n^2+4n+1)-(4n^2-4n+1)=8n，n为整数，故8n是8的倍数。",
                "证明：原式=(4n^2+4n+1)-(4n^2-4n+1)=8n，因为n为整数，所以8n是8的倍数。"
                )["verdict"]
    log("数字序列撞车的成段文字证明题 → undecidable（曾被抢判 numeric_match）",
        v == "undecidable", v)

    # 反过来也要守住：不能为了修上面那条把数值比对一起废掉。
    log("真正的短数值答案仍判 numeric_match",
        cmp_sol("12.50", "12.5")["verdict"] == "numeric_match")
    log("带选项的多段数值答案仍判 numeric_match",
        cmp_sol("①③;12", "①③;12.0")["verdict"] == "numeric_match")

    # 已记录的限制（**未修**，用检查钉住，免得以后被误当成已修）：
    # normalize_answer 不做分数求值，所以 0.5 与 1/2 会被判 mismatch。
    log("[已知限制] 分数与小数不等价：0.5 vs 1/2 仍判 mismatch",
        cmp_sol("0.5", "1/2")["verdict"] == "mismatch",
        "这是记录在案的遗留限制，不是回归")

    # ------------------------------------------------------ 工具错误处理契约
    section("M. 工具错误处理契约（回归：异常抛穿 MCP，可读提示被碾掉）")
    # 回归背景（2026-09-24 真实导出时暴露）：
    #   server.py 的 16 个工具里，11 个没有顶层 except —— 只有为 store.close()
    #   写的 try/finally，而 finally 是不吞异常的。异常抛到 MCP 层会被包成
    #   `UnexpectedToolError: Error executing tool <name>`，工具自己精心写好的
    #   提示全部丢失。最典型的是 zx_export_wrongbook(fmt="xlsx")：
    #   requirements.txt 明写「没装 openpyxl 也会给出明确提示」，
    #   core/export.py 也确实抛了「未安装 openpyxl，请执行 pip install openpyxl」，
    #   但这句话到不了宿主手里。
    #   修法：加一层 @_guard，任何工具失败都返回同一形状的 JSON。
    import asyncio
    import inspect as _inspect

    import server as _srv

    def _text(res):
        for blk in getattr(res, "content", []) or []:
            if getattr(blk, "type", "") == "text":
                return blk.text
        return str(res)

    def _call(name, args):
        return _text(asyncio.run(_srv.mcp.call_tool(name, args)))

    _tools = asyncio.run(_srv.mcp.list_tools())
    _names = [t.name for t in _tools]

    # 1) 每个工具都真的被 _guard 包住了
    _unwrapped = [n for n in _names
                  if not hasattr(getattr(_srv, n, None), "__wrapped__")]
    log(f"全部 {len(_names)} 个工具都有 _guard 兜底", not _unwrapped, str(_unwrapped))

    # 2) functools.wraps 必须把签名转发出去，否则宿主看到的参数全丢
    _sa = next((t for t in _tools if t.name == "submit_analysis"), None)
    _props = sorted((_sa.input_schema.get("properties") or {}).keys()) if _sa else []
    log("兜底层没破坏参数 schema（submit_analysis 8 个参数仍在）",
        len(_props) == 8, str(_props))

    # 3) 参数校验类错误：非法 fmt
    r = json.loads(_call("zx_export_wrongbook", {"fmt": "doc"}))
    log("非法 fmt → 结构化错误而不是抛穿",
        r.get("ok") is False and "fmt" in str(r.get("error", "")), str(r)[:120])

    # 4) 文件不存在
    r = json.loads(_call("zx_import_export_file",
                         {"path": str(TMP / "并不存在.txt"),
                          "subject": "数学", "exam_name": "x"}))
    log("导入不存在的文件 → 结构化错误", r.get("ok") is False, str(r)[:110])

    # 5) 关键：人为注入一个「工具内部未预期异常」，确认也被兜住。
    #    挑 get_questions —— 它的 _store() 调用在 try 之外，属于裸奔路径。
    _orig_store = _srv._store

    def _boom():
        raise RuntimeError("人为注入的 _store 故障")

    _srv._store = _boom
    try:
        r = json.loads(_call("get_questions", {}))
        log("工具内部未预期异常也被兜成结构化错误（get_questions）",
            r.get("ok") is False and r.get("tool") == "get_questions"
            and "人为注入" in str(r.get("error", "")), str(r)[:150])
    except Exception as exc:
        log("工具内部未预期异常也被兜成结构化错误", False,
            f"仍然抛穿：{type(exc).__name__}: {exc}")
    finally:
        _srv._store = _orig_store

    # 6) openpyxl 缺失时那句人话提示，必须能到宿主手里
    import builtins
    _real_import = builtins.__import__

    def _fake_import(name, *a, **kw):
        if name == "openpyxl":
            raise ImportError("人为屏蔽 openpyxl")
        return _real_import(name, *a, **kw)

    builtins.__import__ = _fake_import
    try:
        r = json.loads(_call("zx_export_wrongbook", {"fmt": "xlsx"}))
        log("xlsx 缺 openpyxl 时人话提示能到达宿主（曾被 UnexpectedToolError 碾掉）",
            r.get("ok") is False and "openpyxl" in str(r.get("error", ""))
            and "pip install" in str(r.get("error", "")), str(r.get("error"))[:130])
    finally:
        builtins.__import__ = _real_import

    # ------------------------------------------------------ 练习卷字段类型容忍
    section("N. 练习卷渲染的字段类型容忍（回归：静默拆字 / 直接崩）")
    # 回归背景（2026-09-24）：zx_export_paper 收的是宿主手写的 JSON，
    # 而同一个字段在本项目别的工具里就有两种合法写法 ——
    # submit_analysis 明文接受「| 分隔字符串 或 JSON 数组」。
    # 原来的 render_paper 有两处隐式假设，都会在真实使用中踩到：
    #   · verification 只认 list[dict]；传 submit_solution 的 verdict 字符串
    #     会被逐字符遍历，在 'm'.get(...) 上抛 AttributeError，整个导出失败。
    #   · kp 只认 list；传 "A|B" 会被 '、'.join 逐字拆开成「物、理、/、声…」，
    #     **不报错**，安静地渲染出一份看起来有内容、实际已经花掉的练习卷。
    from core.export import render_paper as _rp

    _base = {"gen_id": "t", "qtype": "单选题", "difficulty": 1.0,
             "stem_html": "<p>x</p>", "answer": "B", "analysis": "",
             "source_topic": "t", "verified": True}

    # 1) verification 传 verdict 字符串
    try:
        h = _rp([{**_base, "kp": ["物理/声与光/光的折射"], "verification": "match"}])
        log("verification 传 verdict 字符串 → 不崩且渲染出记号",
            "✔" in h and "一致" in h)
    except Exception as exc:
        log("verification 传 verdict 字符串 → 不崩", False,
            f"{type(exc).__name__}: {exc}")

    # 2) verification 传 list[str]
    try:
        h = _rp([{**_base, "kp": ["a"], "verification": ["match", "mismatch"]}])
        log("verification 传 list[str] → 不崩", "✔" in h and "✘" in h)
    except Exception as exc:
        log("verification 传 list[str] → 不崩", False, f"{type(exc).__name__}: {exc}")

    # 3) 原有形状（check_practice 的 checks）不能被改坏
    h = _rp([{**_base, "kp": ["a"],
              "verification": [{"name": "题型一致", "kind": "hard",
                                "passed": True, "value": "单选题"}]}])
    log("verification 传 checks（list[dict]）→ 仍渲染明细",
        "题型一致" in h and "✔" in h)

    # 4) kp 传 '|' 分隔字符串 —— 必须拆成两项，不能逐字插顿号
    h = _rp([{**_base, "kp": "物理/声与光/光的折射|物理/声与光/光的反射",
              "verification": []}])
    log("kp 传 '|' 分隔字符串 → 正确拆分（曾被逐字拆成「物、理、/、声…」）",
        "物理/声与光/光的折射、物理/声与光/光的反射" in h and "、物、理" not in h)

    # 5) knowledge_points（不是 kp）传字符串也要认
    _it = {**_base, "knowledge_points": "数学/整式乘除/乘法公式|数学/整式乘除/因式分解",
           "verification": []}
    _it.setdefault("kp", _it.get("knowledge_points", []))   # server.py 里的同一行
    log("knowledge_points 传 '|' 分隔字符串 → 同样正确拆分",
        "数学/整式乘除/乘法公式、数学/整式乘除/因式分解" in _rp([_it]))

    # 6) kp 传 JSON 数组字符串（宿主把列表 json.dumps 了）
    log("kp 传 JSON 数组字符串 → 也能解析",
        "数学/实数/平方根与立方根" in _rp([{**_base, "kp": '["数学/实数/平方根与立方根"]',
                                           "verification": []}]))

    # 7) 空值不炸
    try:
        _rp([{**_base, "kp": None, "verification": None}])
        log("kp / verification 传 None → 不崩", True)
    except Exception as exc:
        log("kp / verification 传 None → 不崩", False, f"{type(exc).__name__}: {exc}")

    # ------------------------------------------------------------ PDF 转换器探测
    section("O. PDF 转换器探测（回归：本机装了浏览器却报「未检测到」）")
    # 回归背景（2026-09-24 真实导出时暴露）：try_pdf 原来只用
    # shutil.which("chrome") / shutil.which("msedge") 探测。
    # Windows 上这两个浏览器**不注册到 PATH** —— 只待在安装目录里，
    # 外加注册表 App Paths（那是给 ShellExecute 用的，which 不查）。
    # 结果本机明明装了 Edge 和 Chrome（标准路径都在），却被报「未检测到」，
    # 一条本来完全能走通的 PDF 导出路径被堵死。
    from core.export import _find_browser, try_pdf as _try_pdf

    _found = _find_browser()
    log("_find_browser 能找到本机浏览器（PATH 查不到时回退到安装目录）",
        _found is not None, str(_found))

    # 关键回归：把 ProgramFiles / ProgramFiles(x86) 抹掉再找一次。
    # 这正是 2026-09-24 真实踩到的那一层 —— 某些宿主/沙箱里这两个环境变量
    # 直接是 None（PATH 也被洗过），只靠 os.environ 就会漏掉明明装在
    # 标准位置的浏览器。所以必须靠字面量路径兜底。
    _saved = {k: os.environ.pop(k, None)
              for k in ("ProgramFiles", "ProgramFiles(x86)")}
    try:
        _found2 = _find_browser()
        log("环境变量被抹掉后仍能找到浏览器（靠字面量标准路径兜底）",
            _found2 is not None or _found is None, str(_found2))
    finally:
        for _k, _v in _saved.items():
            if _v is not None:
                os.environ[_k] = _v

    _h = TMP / "pdf_probe.html"
    _h.write_text("<html><body><h1>PDF 探针</h1><p>中文测试</p></body></html>",
                  encoding="utf-8")
    _r = _try_pdf(_h)
    _pdf = Path(_r["pdf"]) if _r.get("pdf") else None
    _real = bool(_r.get("available") is True and _pdf is not None
                 and _pdf.exists() and _pdf.stat().st_size > 0)
    if _found is not None:
        # 探测到了就必须真出文件 —— 不能只信 available=True
        log("try_pdf 真产出 PDF（不只是报告 available=True）", _real,
            f"converter={_r.get('converter')}，"
            f"size={_pdf.stat().st_size if _real else '-'} B")
    else:
        log("本机确实没有浏览器 → try_pdf 明确说没有而不是假装成功",
            _r.get("available") is False and "how_to" in _r,
            str(_r.get("reason"))[:90])

    # ==================================================================== P
    section("P. 派生字段覆盖规则（回归：重新同步静默清空全部分析结果）")
    # 回归背景（2026-09-25，真实账号实测，P0 级）：
    #   core/store.py 的 upsert() 走的是**无条件全字段 UPDATE**，
    #   而 core/models.py 里 `analysis = self.analysis.model_dump() if self.analysis else None`
    #   —— 新同步抓回来的题天然没有分析结果（分析是宿主后来提交的），
    #   于是这个 None 直接盖掉了库里已有的分析。
    #
    #   实测后果：27 道题的 analysis 从 27 条掉到 0 条，
    #   而 zx_sync 返回的是 ok=true、updated=27，**一句警告都没有**。
    #   毁掉的偏偏是最贵的东西（AI 花 token 读题 + 读手写图产出的结论）。
    #
    # ⚠️ 为什么原来的测试没抓到：上面第 I 段有「二次同步零新增」，
    #   但它**只断言了题数不变**（不会变多），没有断言「分析结果还在」（不会变少）。
    #   所以这一段的断言刻意做成**双向**：不会变多、不会变少、内容也不变。
    from core.models import Analysis as _An
    from core.models import AnswerPart as _AP3
    from core.models import Exam as _Ex3
    from core.models import HistoryEntry as _HE
    from core.models import PracticeRef as _PR
    from core.models import QuestionPart as _QP3
    from core.models import Score as _Sc3
    from core.models import WrongQuestion as _WQ3

    dstore = Store(TMP / "derive.db", TMP / "images_derive")

    def _dq():
        """同一道题的新抓取对象 —— 注意它**没有**派生字段。"""
        return _WQ3(id="d1", source="manual", subject="物理",
                    exam=_Ex3(name="派生字段回归", date="2026-09-25"),
                    question=_QP3(stem_html="<p>派生字段回归题干</p>"),
                    answer=_AP3(standard="42"), score=_Sc3())

    fp_d = _dq().fingerprint()
    dstore.upsert(_dq())

    # 宿主后来提交了分析结果 + 练习记录 + 历史
    q1 = dstore.get(fp_d)
    q1.analysis = _An(error_type="计算错误",
                      knowledge_points=["物理/电与磁/欧姆定律"],
                      evidence=["42"], confidence=0.8,
                      analyzed_by="selftest", prompt_version="analyze-v3")
    q1.practice = [_PR(gen_id="p_1", verified=True)]
    q1.history = [_HE(date="2026-09-20", result="correct")]
    dstore.upsert(q1)

    before = dstore.get(fp_d)
    log("前置条件：分析/练习/历史确实写进去了",
        before.analysis is not None and len(before.practice) == 1
        and len(before.history) == 1, "")

    # 模拟「重新同步」：同一道题、内容相同，但新对象不带派生字段
    dstore.upsert(_dq())
    after = dstore.get(fp_d)

    log("重新同步后 analysis 仍在（回归：27 条分析被清成 0 条）",
        after.analysis is not None and after.analysis.error_type == "计算错误",
        f"analysis={'保留' if after.analysis else '❌ 没了'}")
    log("重新同步后 practice / history 仍在",
        len(after.practice) == 1 and len(after.history) == 1,
        f"practice={len(after.practice)} history={len(after.history)}")
    log("重新同步后分析结果**内容**一字未改",
        after.analysis.model_dump() == before.analysis.model_dump(), "")
    log("题数没变多（原有的去重承诺仍然成立）",
        dstore.counts()["total"] == 1, f"total={dstore.counts()['total']}")

    # 反向：带了新分析结果时必须照常更新 —— 别修过头变成「永远不更新」
    q2 = dstore.get(fp_d)
    q2.analysis = _An(error_type="审题失误",
                      knowledge_points=["物理/电与磁/欧姆定律"],
                      evidence=["42"], confidence=0.9,
                      analyzed_by="selftest", prompt_version="analyze-v3")
    dstore.upsert(q2)
    log("带了新分析结果时照常更新（没修成「永不更新」）",
        dstore.get(fp_d).analysis.error_type == "审题失误",
        dstore.get(fp_d).analysis.error_type)

    log("「保留了多少条既有分析」是可观察的（不是隐性承诺）",
        dstore.preserved_analysis >= 1,
        f"preserved_analysis={dstore.preserved_analysis}")
    dstore.close()

    # ==================================================================== Q
    section("Q. 参数形状与未知参数（回归：ok=true 但产物不存在 / 传数组被拒）")

    # -- Q1 未知参数必须被拒绝，而不是静默忽略
    #
    # 回归背景（2026-09-25 实测）：pydantic 默认 extra='ignore'，于是
    #   zx_export_paper(items=..., fmt="pdf")   → ok=true，但没有 PDF
    #   zx_export_paper(items=..., 随便写="x")  → ok=true，静默通过
    # `fmt` 是兄弟工具 zx_export_wrongbook 的参数，宿主很容易顺手套过来。
    # 这是最危险的一类失败：不报错，结果是错的。
    import asyncio as _aio
    import server as _srv

    log("未知参数拦截器已安装（库换了漏斗时这条会失败，而不是悄悄失效）",
        bool(_srv.UNKNOWN_ARG_GUARD.get("installed")),
        str(_srv.UNKNOWN_ARG_GUARD))

    def _txt(res):
        for b in getattr(res, "content", []) or []:
            if getattr(b, "type", "") == "text":
                return b.text
        return str(res)

    async def _call(tool, payload):
        return json.loads(_txt(await _srv.mcp.call_tool(tool, payload)))

    _items = json.dumps([{"gen_id": "q1", "qtype": "解答题",
                          "knowledge_points": ["物理/电与磁/欧姆定律"],
                          "difficulty": 0.5, "stem_html": "<p>形状测试</p>",
                          "answer": "2A"}], ensure_ascii=False)

    r_unknown = _aio.run(_call("zx_export_paper",
                               {"items": _items, "pdf": True, "fmt": "pdf"}))
    log("传未知参数 pdf → ok=false 并点名是哪个参数（不再静默成功）",
        r_unknown.get("ok") is False and "pdf" in str(r_unknown.get("error")),
        str(r_unknown.get("error"))[:70])

    r_bogus = _aio.run(_call("zx_export_paper",
                             {"items": _items, "完全不存在": "x"}))
    log("传完全瞎写的参数名 → 同样被拒绝",
        r_bogus.get("ok") is False, str(r_bogus.get("error"))[:70])

    r_ok = _aio.run(_call("zx_export_paper",
                          {"items": _items, "title": "形状回归", "want_pdf": False}))
    log("正确参数照常工作（没把工具卡死）",
        r_ok.get("ok") is True, str(r_ok.get("html"))[-40:])

    # fmt 兼容别名：宿主套用兄弟工具（zx_export_wrongbook）的习惯时应该也能用
    log("zx_export_paper 认下 fmt='pdf' 这个兼容写法（并真的走 PDF 分支）",
        _aio.run(_call("zx_export_paper",
                       {"items": _items, "fmt": "pdf"})).get("ok") is True, "")

    # -- Q2 列表型参数要接受数组（原来只认 str，传数组在调用层就被拒）
    from server import _as_list as _al

    log("_as_list：数组原样", _al(["a", "b"]) == ["a", "b"], "")
    log("_as_list：JSON 数组字符串", _al('["a","b"]') == ["a", "b"], "")
    log("_as_list：竖线分隔", _al("a|b") == ["a", "b"], "")
    log("_as_list：逗号分隔（zx_sync 用）", _al("数学,物理", ",") == ["数学", "物理"], "")
    log("_as_list：None / 空串 → 空表", _al(None) == [] and _al("  ") == [], "")
    # 这条是刻意保护：evidence 里可能引用含逗号的原文，
    # 若按逗号切分会把一条证据切成两条 —— 那是新的静默错误
    log("_as_list：默认不按逗号切（保护含逗号的证据原文）",
        _al("学生把 3,4 看成了 3.4") == ["学生把 3,4 看成了 3.4"], "")

    # 端到端：往 server 真正在用的那个库里塞一道题，再拿它的指纹提交分析
    # （server 用的是 ZX_DB_PATH 指向的 edge.db，不能用别的库的指纹）
    _wq_edge = _dq()
    store.upsert(_wq_edge)
    _fp_edge = _wq_edge.fingerprint()

    r_arr = _aio.run(_call("submit_analysis", {
        "fingerprint": _fp_edge, "error_type": "概念不清",
        "knowledge_points": ["物理/电路与欧姆定律/欧姆定律"],
        "evidence": ["42"], "confidence": 0.5}))
    log("submit_analysis 接受**数组**形状的 knowledge_points / evidence（不再 ToolError）",
        r_arr.get("ok") is True,
        f"ok={r_arr.get('ok')} rejected={r_arr.get('rejected')} "
        f"errors={str(r_arr.get('errors'))[:60]}")

    r_sync_arr = _aio.run(_call("zx_sync", {"subjects": ["物理"], "max_exams": 0}))
    log("zx_sync 接受**数组**形状的 subjects（不再 ToolError）",
        isinstance(r_sync_arr, dict) and "ToolError" not in str(r_sync_arr),
        str(r_sync_arr.get("error"))[:60])

    # ==================================================================== R
    section("R. 会话状态三态（回归：Cookie 过期时漏出 TypeError）")
    # 回归背景（2026-09-25 实测）：zx_session_status 原来直接 try: login()，
    # 而库在会话失效时抛的是 `TypeError: 'NoneType' object is not subscriptable`
    # （zhixuewang/account.py:18 —— 服务端返回 result=null，库直接对它取 ["role"]）。
    # 使用者分不清「Cookie 过期」「网络不通」「代码坏了」，而这三件事该做的事完全不同。
    #
    # 这里用**假响应**离线验证分类逻辑，不发真实请求。
    from adapters import zhixue_web as _zw

    class _FakeResp:
        def __init__(self, payload=None, status=200, text="", boom=False):
            self._p, self.status_code, self.ok, self.text = payload, status, status == 200, text
            self._boom = boom

        def json(self):
            if self._boom:
                raise ValueError("not json")
            return self._p

    _real_get = _zw.requests.get
    try:
        # 会话有效：result 是有内容的 dict（注意成功码可能是 200 也可能是 0）
        _zw.requests.get = lambda *a, **k: _FakeResp(
            {"errorCode": 200, "errorInfo": "操作成功",
             "result": {"id": "1", "name": "某同学", "role": "student"}})
        _r = _zw.session_check("loginUserName=u; token=t")
        log("会话有效 → status=valid（成功码 200 也认，不只看 errorCode=0）",
            _r.get("status") == "valid" and _r.get("role") == "student", str(_r))

        _zw.requests.get = lambda *a, **k: _FakeResp(
            {"errorCode": 0, "errorInfo": "操作成功",
             "result": {"id": "1", "name": "某同学", "role": "student"}})
        log("会话有效 → status=valid（成功码 0 也认）",
            _zw.session_check("loginUserName=u; token=t").get("status") == "valid", "")

        # 会话失效：result=null —— 这正是库抛 TypeError 的那种响应
        _zw.requests.get = lambda *a, **k: _FakeResp(
            {"errorCode": -200, "errorInfo": "操作失败", "result": None})
        _r = _zw.session_check("loginUserName=u; token=t")
        log("会话失效 → status=expired，并带出 errorCode（而不是 TypeError）",
            _r.get("status") == "expired" and _r.get("error_code") == -200, str(_r))

        # 网络不通：不能报成「Cookie 失效」，否则会让人白重新登录一次
        def _boom(*a, **k):
            raise ConnectionError("代理拒绝连接")
        _zw.requests.get = _boom
        _r = _zw.session_check("loginUserName=u; token=t")
        log("网络不通 → status=unreachable（明确说「这不代表 Cookie 失效」）",
            _r.get("status") == "unreachable" and "不代表" in str(_r.get("note")), str(_r))

        # 响应不是 JSON
        _zw.requests.get = lambda *a, **k: _FakeResp(status=502, text="<html>502</html>")
        _r = _zw.session_check("loginUserName=u; token=t")
        log("非 200 响应 → unreachable，不误判成过期",
            _r.get("status") == "unreachable", str(_r))
    finally:
        _zw.requests.get = _real_get

    # 端到端：伪造一个失效 Cookie，确认 zx_session_status 不再漏 TypeError
    # （走独立的凭据命名空间，物理上碰不到真实凭据）
    import importlib
    from adapters import session as _sess_mod

    _env_bak = {k: os.environ.get(k) for k in
                ("ZX_CRED_SERVICE", "ZX_CRED_KEY", "ZX_CRED_FILE")}
    os.environ["ZX_CRED_SERVICE"] = "zhixue-wrongbook-edge-sess"
    os.environ["ZX_CRED_KEY"] = "zx_cookie_edge_sess"
    os.environ["ZX_CRED_FILE"] = str(TMP / ".session_edge_sess.bin")
    try:
        importlib.reload(_sess_mod)
        _sess_mod.set_cookie(
            "loginUserName=abcdefg; tlsysSessionId=deadbeefdeadbeefdeadbeef; "
            "deviceId=00000000-0000-0000-0000-000000000000")
        # 注意 session.status() 返回的是 dict（JSON 序列化在 server.py 那层做）
        _st = _sess_mod.status()
        log("伪造 Cookie 能存进独立命名空间（前置条件）",
            _st.get("has_cookie") is True, "")
        log("确认没碰到真实凭据命名空间",
            _sess_mod.SERVICE_NAME == "zhixue-wrongbook-edge-sess",
            _sess_mod.SERVICE_NAME)
        _sess_mod.delete_cookie()
        log("清理：独立命名空间里的假 Cookie 已删除",
            _sess_mod.status().get("has_cookie") is False, "")
    finally:
        for k, v in _env_bak.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        importlib.reload(_sess_mod)

    # ==================================================================== S
    section("S. 数据路径解析（回归：环境变量传 Git Bash 路径 → 静默建空库）")
    # 回归背景（2026-09-25 实测踩到）：在 Git Bash 里用 $PWD 给 ZX_DB_PATH 赋值，
    # 得到的是 `/d/Data/...` 这种 POSIX 风格路径。Python 在 Windows 上把它当成
    # 「当前盘符下的相对路径」`\d\Data\...` —— 于是**静默新建了一个空库**，
    # 工具照常运行，只是所有科目都显示「0 道题」，
    # 使用者看到「本地还没有任何错题数据」，完全猜不到是路径错了。
    from core.config import get_config as _get_cfg

    _env_bak = os.environ.get("ZX_DB_PATH")
    try:
        if sys.platform == "win32":
            os.environ["ZX_DB_PATH"] = "/d/Data/Desktop/x.db"
            _get_cfg.cache_clear()
            log("Git Bash 风格路径自动转换（不再静默建空库）",
                str(_get_cfg().path("db_path")) == "D:\\Data\\Desktop\\x.db",
                str(_get_cfg().path("db_path")))

        os.environ["ZX_DB_PATH"] = "rel/y.db"
        _get_cfg.cache_clear()
        _p = _get_cfg().path("db_path")
        log("相对路径按 ROOT 解析（不受当前工作目录影响）",
            _p.is_absolute() and _p.parent.parent.name == "zhixue-wrongbook",
            str(_p))

        os.environ["ZX_DB_PATH"] = str(TMP / "abs.db")
        _get_cfg.cache_clear()
        log("绝对路径原样使用",
            str(_get_cfg().path("db_path")) == str(TMP / "abs.db"),
            str(_get_cfg().path("db_path")))
    finally:
        if _env_bak is None:
            os.environ.pop("ZX_DB_PATH", None)
        else:
            os.environ["ZX_DB_PATH"] = _env_bak
        _get_cfg.cache_clear()

    section("T. 诊断范围闸门（回归：只问科目、不问「分析哪些试卷」就出报告）")
    # 回归背景（2026-09-25 新增需求）：原来只强制问「哪一科」。
    #   但一科的错题可能横跨几个月、几十份卷子（周测 / 午练 / 晚练 / 早读），
    #   只说「物理」的话，「分析哪些卷子」仍然是**模型替用户定的** ——
    #   范围没被用户确认，等于没确认。现在必须两问都答：
    #   ① 科目  ② 试卷范围（类型 × 时间，两个维度是「且」）
    #
    # ⚠️ 这段刻意测三层，而不只测「有没有拒绝」：
    #   1) 闸门的七种状态 —— 拒绝理由要能区分，否则宿主不知道下一步问什么；
    #   2) 范围**真的过滤了数据** —— 否则闸门只是句问话，
    #      用户选了「本周午练」却拿到全量结论，比不问更糟；
    #   3) 审计日志记下范围 —— 事后能核对「这次到底覆盖了什么」。
    from datetime import date as _date
    from datetime import timedelta as _td

    from core.constants import TIME_SCOPES as _TS
    from core.profile import build_profile as _bp
    from core.scope import (normalize_paper_types as _npt)
    from core.scope import (normalize_time_scope as _nts)
    from core.scope import (paper_type_options as _pto)
    from core.scope import (resolve_scope as _rs)
    from core.scope import (resolve_time_scope as _rts)
    from core.scope import (time_scope_options as _tso)
    from core.scope import (valid_paper_types as _vpt)
    from core.store import paper_type_where as _ptw
    from server import _diagnosis_gate as _gate

    # -- 归一化：用户/宿主怎么说都得能懂 ---------------------------------
    log("类型归一化：单值 / 逗号 / 竖线 / 数组",
        _npt("周测") == ["周测"] and _npt("周测,午练") == ["周测", "午练"]
        and _npt("午练|晚练") == ["午练", "晚练"]
        and _npt(["晚练", "早读"]) == ["晚练", "早读"], "")
    log("类型归一化：别名与「不限」",
        _npt("周考") == ["周测"] and _npt("不限") == ["全部"]
        and _npt("其它") == ["其他"], "")
    # 这条是刻意保护：不认识的词**原样保留**，交给闸门报错。
    # 静默丢弃会让「传错了」看起来像「筛过了」—— 那是最危险的一类 bug。
    log("类型归一化：不认识的词原样保留（不静默丢弃）",
        _npt("火星卷") == ["火星卷"], str(_npt("火星卷")))
    log("时间归一化：这周→本周 / 最近两周→近两周 / 不限→全部",
        _nts("这周") == "本周" and _nts("最近两周") == "近两周"
        and _nts("不限") == "全部", "")

    # -- 时间区间解析（用固定 today，保证可复现）--------------------------
    _today = _date(2026, 9, 25)          # 周五
    log("本周 = 本周一 ~ 今天（周一起算，不是周日）",
        _rts("本周", today=_today)["since"] == "2026-09-21"
        and _rts("本周", today=_today)["until"] == "2026-09-25",
        str(_rts("本周", today=_today)["label"]))
    log("上周 = 上一个完整周（周一~周日）",
        _rts("上周", today=_today)["since"] == "2026-09-14"
        and _rts("上周", today=_today)["until"] == "2026-09-20",
        str(_rts("上周", today=_today)["label"]))
    log("近两周 / 本月 / 近一月 起止正确",
        _rts("近两周", today=_today)["since"] == "2026-09-12"
        and _rts("本月", today=_today)["since"] == "2026-09-01"
        and _rts("近一月", today=_today)["since"] == "2026-08-27", "")
    log("「全部」= 不按时间筛（起止为空）",
        _rts("全部", today=_today)["since"] == ""
        and _rts("全部", today=_today)["until"] == "", "")
    log("自定义区间：2026/9/1 与 20260901 都能认",
        _rts("自定义", since="2026/9/1", until="2026-09-10")["since"] == "2026-09-01"
        and _rts("自定义", since="20260901", until="")["since"] == "2026-09-01", "")

    # -- SQL 条件：这里最容易写错，所以单独测 ----------------------------
    log("「全部」不产生任何过滤条件",
        _ptw(["全部"]) is None and _ptw(None) is None and _ptw([]) is None, "")
    _sql_other = _ptw(["其他"])
    # 「其他」必须是 4 个 NOT LIKE 的 **AND**。写成 OR 会永远成立（等于没筛），
    # 而且不会报错 —— 只会在数据上表现为「其他 = 全量」。
    log("「其他」= 4 个 NOT LIKE 的 AND（不是 OR）",
        _sql_other is not None and _sql_other[0].count("AND") == 3
        and _sql_other[0].count("OR") == 0 and len(_sql_other[1]) == 4,
        _sql_other[0][:70] if _sql_other else "-")
    log("多选类型 = OR 关系",
        _ptw(["周测", "午练"]) is not None and "OR" in _ptw(["周测", "午练"])[0], "")

    # -- 造一批跨类型、跨时间的题 ----------------------------------------
    tstore = Store(TMP / "scope.db", TMP / "images_scope")

    def _sq(exam_name, days_ago, stem):
        return _WQ3(id=stem, source="manual", subject="物理",
                    exam=_Ex3(name=exam_name,
                              date=(_today - _td(days=days_ago)).isoformat()),
                    question=_QP3(stem_html=f"<p>{stem}</p>"),
                    answer=_AP3(standard="42"), score=_Sc3())

    tstore.upsert(_sq("2026秋周测A", 0, "周测题"))
    tstore.upsert(_sq("午练-八年级-20260922", 3, "午练题"))
    tstore.upsert(_sq("晚练-八年级-20260901", 40, "晚练题"))
    tstore.upsert(_sq("2026秋期末测试", 5, "期末题"))

    _cnt = tstore.counts_by_paper_type()
    log("按类型计数正确（周测/午练/晚练/其他 各 1，早读 0）",
        _cnt["周测"] == 1 and _cnt["午练"] == 1 and _cnt["晚练"] == 1
        and _cnt["早读"] == 0 and _cnt["其他"] == 1 and _cnt["__total__"] == 4,
        str({k: v for k, v in _cnt.items() if k != "__total__"}))
    # 这条是「其他」写错成 OR 时的兜底断言：一旦写错，其他 = 4（全量）
    log("「其他」只命中不属于四类的那 1 道（写错成 OR 就会变 4）",
        tstore.count_questions(paper_types=["其他"]) == 1, "")
    log("时间过滤生效：近一月排除 40 天前那道",
        tstore.count_questions(date_from="2026-08-27", date_to="2026-09-25") == 3, "")
    log("类型 × 时间 是「且」关系",
        tstore.count_questions(paper_types=["周测"], date_from="2026-09-21") == 1
        and tstore.count_questions(paper_types=["晚练"], date_from="2026-09-21") == 0,
        "")

    _opts = _pto(tstore, "物理")
    log("选项清单带上每题数（宿主才能照着问）",
        _opts["total"] == 4
        and {r["type"]: r["questions"] for r in _opts["paper_types"]}["其他"] == 1,
        str(_opts["with_questions"]))
    log("时间预设清单覆盖全部预设",
        len(_tso(tstore, "物理", today=_today)["time_scopes"]) == len(_TS), "")

    # -- 闸门七种状态 -----------------------------------------------------
    def _flag(d):
        return [k for k in ("needs_subject", "needs_confirmation", "needs_scope",
                            "invalid_scope", "needs_scope_confirmation",
                            "no_data", "no_data_in_scope") if d and d.get(k)]

    log("① 什么都不给 → needs_subject",
        _flag(_gate(tstore, "", False)) == ["needs_subject"], "")
    log("② 给科目未确认 → needs_confirmation",
        _flag(_gate(tstore, "物理", False)) == ["needs_confirmation"], "")
    log("③ 确认科目但不给范围 → needs_scope（本次新增的核心）",
        _flag(_gate(tstore, "物理", True)) == ["needs_scope"], "")
    log("③ 拒绝时带回可选的类型与时间清单（宿主照着问）",
        bool(_gate(tstore, "物理", True).get("paper_types"))
        and bool(_gate(tstore, "物理", True).get("time_scopes")), "")
    log("④ 范围取值非法 → invalid_scope（不静默忽略）",
        _flag(_gate(tstore, "物理", True, paper_types=["火星卷"],
                    time_scope="本周")) == ["invalid_scope"], "")
    log("④ 非法时告诉宿主合法取值有哪些",
        _gate(tstore, "物理", True, paper_types=["火星卷"],
              time_scope="本周").get("valid_paper_types") == _vpt(), "")
    log("⑤ 范围未确认 → needs_scope_confirmation",
        _flag(_gate(tstore, "物理", True, paper_types=["周测"],
                    time_scope="全部")) == ["needs_scope_confirmation"], "")
    log("⑥ 科目没数据 → no_data",
        _flag(_gate(tstore, "地理", True, paper_types=["全部"],
                    time_scope="全部", scope_confirmed=True)) == ["no_data"], "")
    log("⑦ 范围里没题 → no_data_in_scope（不是悄悄放宽成全部）",
        _flag(_gate(tstore, "物理", True, paper_types=["早读"],
                    time_scope="全部", scope_confirmed=True))
        == ["no_data_in_scope"], "")
    log("⑦ 拒绝时提示「别直接放宽成全部，让用户重选」",
        "放宽" in str(_gate(tstore, "物理", True, paper_types=["早读"],
                            time_scope="全部",
                            scope_confirmed=True).get("hint", "")), "")
    log("⑧ 两问都答 → 放行（返回 None）",
        _gate(tstore, "物理", True, paper_types=["周测"], time_scope="全部",
              scope_confirmed=True) is None, "")

    # -- 最关键的一条：范围必须真的过滤数据 -------------------------------
    _p_all = _bp(tstore, subject="物理")
    _p_week = _bp(tstore, subject="物理", paper_types=["周测"])
    # 09-20 之后的三道（周测 09-25 / 午练 09-22 / 期末 09-20），
    # 40 天前那道晚练（08-16）被排除 —— 注意期末题是 09-20，别把下界写成 09-21
    _p_time = _bp(tstore, subject="物理", date_from="2026-09-20")
    log("画像真的被类型过滤（全部 4 道 → 只周测 1 道）",
        _p_all["question_count"] == 4 and _p_week["question_count"] == 1,
        f"全部={_p_all['question_count']} 周测={_p_week['question_count']}")
    log("画像真的被时间过滤（40 天前那道被排除）",
        _p_time["question_count"] == 3, str(_p_time["question_count"]))
    log("画像回显了本次覆盖范围（报告里要写给用户看）",
        _p_week["scope"]["paper_types"] == ["周测"], str(_p_week["scope"]))
    log("范围标签可读（人能看懂「周测 + 上周」这种）",
        "周测" in _rs(["周测"], "上周", today=_today)["label"]
        and "上周" in _rs(["周测"], "上周", today=_today)["label"],
        _rs(["周测"], "上周", today=_today)["label"])

    # -- 审计日志必须记下范围（不留痕的规则等于没有规则）------------------
    _aid = tstore.log_diagnosis(subject="物理", user_confirmed=True,
                                question_count=1, analyzed_count=1,
                                weak_top=["物理/电路与欧姆定律/欧姆定律"],
                                paper_types=["周测", "午练"],
                                time_scope="本周", scope_confirmed=True)
    _log1 = tstore.diagnosis_logs(limit=1)[0]
    log("审计日志记下范围与两处确认声明",
        _log1["id"] == _aid and _log1["paper_types"] == ["周测", "午练"]
        and _log1["time_scope"] == "本周" and _log1["scope_confirmed"] is True,
        f"types={_log1['paper_types']} scope={_log1['time_scope']}")

    # -- 老库迁移：diagnosis_log 缺列时必须自动补上 ------------------------
    _old = TMP / "oldlog.db"
    _raw = sqlite3.connect(_old)
    _raw.execute("""CREATE TABLE diagnosis_log (
        id INTEGER PRIMARY KEY AUTOINCREMENT, created_at TEXT NOT NULL,
        subject TEXT NOT NULL, user_confirmed INTEGER NOT NULL DEFAULT 0,
        question_count INTEGER DEFAULT 0, analyzed_count INTEGER DEFAULT 0,
        weak_top TEXT, host TEXT)""")
    _raw.commit()
    _raw.close()
    _mig = Store(_old, TMP / "images_oldlog")
    _mcols = {r["name"] for r in _mig.conn.execute(
        "PRAGMA table_info(diagnosis_log)")}
    log("老库自动补列（paper_types / time_scope / scope_confirmed）",
        {"paper_types", "time_scope", "scope_confirmed"} <= _mcols,
        str(sorted(_mcols)))
    # 补列之后必须**真的能写**，否则迁移只是看着对
    _mig.log_diagnosis(subject="物理", user_confirmed=True, paper_types=["周测"],
                       time_scope="上周", scope_confirmed=True)
    log("迁移后的老库能正常写审计（不只是列名对）",
        _mig.diagnosis_logs(limit=1)[0]["time_scope"] == "上周", "")
    _mig.close()
    tstore.close()

    # ------------------------------------------------------------------ 收尾
    store.close()
    s = Store(TMP / "edge.db", TMP / "images")
    s.purge(confirm=True)
    shutil.rmtree(TMP, ignore_errors=True)

    passed = sum(1 for _, ok, _ in RESULTS if ok)
    print("\n" + "=" * 68)
    print(f"汇总：{passed}/{len(RESULTS)} 通过")
    for n, ok, d in RESULTS:
        if not ok:
            print(f"  FAIL {n} — {d}")
    print("=" * 68)
    return 0 if passed == len(RESULTS) else 1


def _build_minimal_pdf(lines: list[str], draw_only: bool = False) -> bytes:
    """手搓一个最小可用的单页 PDF，用来验证 pypdf 抽取通链路。

    刻意只用 ASCII：中文需要嵌入字体，那是另一个话题；
    这里要验的是「PDF → 文本 → 题块」这条管道本身通不通。
    draw_only=True 时只画一个矩形、不放文字，模拟扫描版（无文本层）。
    """
    parts = ["BT", "/F1 12 Tf", "72 720 Td"]
    if draw_only:
        parts = ["0 0 1 rg", "100 100 200 200 re", "f"]
    else:
        for i, ln in enumerate(lines):
            if i:
                parts.append("0 -16 Td")
            esc = ln.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            parts.append(f"({esc}) Tj")
        parts.append("ET")
    stream = "\n".join(parts).encode("latin-1")

    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
        + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    n = len(objs) + 1
    out += f"xref\n0 {n}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {n} /Root 1 0 R >>\nstartxref\n{xref_pos}\n"
            f"%%EOF\n").encode()
    return bytes(out)


if __name__ == "__main__":
    sys.exit(main())
